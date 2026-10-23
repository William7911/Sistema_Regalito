import csv
import io
from datetime import datetime
from typing import Optional, List
from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import select

from app.db.models import Cliente, ClienteContacto, TipoCliente, Municipio
from app.schemas import cliente_schema as schemas
from app.services.audit_service import log_action


class ConcreteClienteService:
    """Lógica de negocio para Clientes y Tipos de Cliente (DERCAS).
    Utiliza UnitOfWork inyectado para gestionar transacciones y repositorios."""

    def __init__(self, uow):
        self.uow = uow

    async def _validate_tipo_cliente(self, id_tipo: int):
        tipo = await self.uow.tipos_cliente.get_by_id(id_tipo)
        if not tipo:
            raise HTTPException(status_code=400, detail="El tipo de cliente especificado no existe")
        if tipo.estado != "Activo":
            raise HTTPException(status_code=400, detail="El tipo de cliente especificado está inactivo")
        return tipo

    async def _validate_municipio(self, id_municipio: Optional[int]):
        if id_municipio is None:
            return
        res = await self.uow.db.execute(select(Municipio).where(Municipio.id_municipio == id_municipio))
        if not res.scalar_one_or_none():
            raise HTTPException(status_code=400, detail="El municipio especificado no existe")

    def _to_out(self, c: Cliente) -> schemas.ClienteOut:
        muni_nombre = c.municipio.nombre if c.municipio else None
        tipo_out = schemas.TipoClienteOut.model_validate(c.tipo_cliente) if c.tipo_cliente else None
        contactos_out = [schemas.ClienteContactoOut.model_validate(ct) for ct in (c.contactos or [])]
        return schemas.ClienteOut(
            id_cliente=c.id_cliente,
            id_tipo_cliente=c.id_tipo_cliente,
            id_municipio=c.id_municipio,
            nit=c.nit,
            nombre_comercial=c.nombre_comercial,
            direccion=c.direccion,
            estado=c.estado,
            fecha_creacion=c.fecha_creacion,
            contactos=contactos_out,
            tipo_cliente=tipo_out,
            municipio_nombre=muni_nombre,
        )

    # -----------------------------------------------------------------------
    # Métodos de Catálogo: Tipo de Cliente
    # -----------------------------------------------------------------------

    async def list_tipos_cliente(self, only_active: bool = False) -> List[schemas.TipoClienteOut]:
        tipos = await self.uow.tipos_cliente.list_all(only_active=only_active)
        return [schemas.TipoClienteOut.model_validate(t) for t in tipos]

    async def list_municipios(self, id_departamento: Optional[int] = None) -> List[schemas.MunicipioOut]:
        stmt = select(Municipio).where(Municipio.estado == "Activo")
        if id_departamento is not None:
            stmt = stmt.where(Municipio.id_departamento == id_departamento)
        stmt = stmt.order_by(Municipio.nombre.asc())
        res = await self.uow.db.execute(stmt)
        return [schemas.MunicipioOut.model_validate(m) for m in res.scalars().all()]

    async def create_tipo_cliente(self, data: schemas.TipoClienteCreate) -> schemas.TipoClienteOut:
        existing = await self.uow.tipos_cliente.get_by_nombre(data.nombre)
        if existing:
            raise HTTPException(status_code=400, detail="Ya existe un tipo de cliente con ese nombre")
        tipo = TipoCliente(
            nombre=data.nombre.strip(),
            descuento_predeterminado=data.descuento_predeterminado,
            estado="Activo",
        )
        created = await self.uow.tipos_cliente.create(tipo)
        await self.uow.commit()
        return schemas.TipoClienteOut.model_validate(created)

    async def get_tipo_cliente(self, id_tipo: int) -> schemas.TipoClienteOut:
        tipo = await self.uow.tipos_cliente.get_by_id(id_tipo)
        if not tipo:
            raise HTTPException(status_code=404, detail="Tipo de cliente no encontrado")
        return schemas.TipoClienteOut.model_validate(tipo)

    async def update_tipo_cliente(
        self, id_tipo: int, data: schemas.TipoClienteUpdate
    ) -> schemas.TipoClienteOut:
        tipo = await self.uow.tipos_cliente.get_by_id(id_tipo)
        if not tipo:
            raise HTTPException(status_code=404, detail="Tipo de cliente no encontrado")

        if data.nombre is not None:
            nombre = data.nombre.strip()
            duplicate = await self.uow.tipos_cliente.get_by_nombre(nombre)
            if duplicate and duplicate.id_tipo_cliente != id_tipo:
                raise HTTPException(status_code=400, detail="Ya existe un tipo de cliente con ese nombre")
            tipo.nombre = nombre
        if data.descuento_predeterminado is not None:
            tipo.descuento_predeterminado = data.descuento_predeterminado
        if data.estado is not None:
            tipo.estado = data.estado

        updated = await self.uow.tipos_cliente.update(tipo)
        await self.uow.commit()
        return schemas.TipoClienteOut.model_validate(updated)

    async def deactivate_tipo_cliente(
        self, id_tipo: int, actor_id: int = 0
    ) -> schemas.TipoClienteOut:
        tipo = await self.uow.tipos_cliente.get_by_id(id_tipo)
        if not tipo:
            raise HTTPException(status_code=404, detail="Tipo de cliente no encontrado")
        tipo.estado = "Inactivo"
        await self.uow.tipos_cliente.update(tipo)
        await self.uow.commit()

        audit_actor = actor_id or 1
        await log_action(
            self.uow.db,
            id_usuario=audit_actor,
            tabla="tipo_cliente",
            accion="DESACTIVAR",
            descripcion=f"Desactivó tipo de cliente '{tipo.nombre}' (ID {tipo.id_tipo_cliente}).",
        )
        await self.uow.commit()
        return schemas.TipoClienteOut.model_validate(tipo)

    # -----------------------------------------------------------------------
    # Métodos de Entidad: Cliente
    # -----------------------------------------------------------------------

    async def list_clientes(self, filters: schemas.ClienteFilter) -> List[schemas.ClienteOut]:
        clientes = await self.uow.clientes.search(
            nombre=filters.nombre,
            nit=filters.nit,
            id_tipo_cliente=filters.id_tipo_cliente,
            is_active=filters.is_active,
            fecha_desde=filters.fecha_desde,
            fecha_hasta=filters.fecha_hasta,
            sort_by=filters.sort_by,
            sort_order=filters.sort_order,
            limit=filters.limit,
            offset=filters.offset,
        )
        return [self._to_out(c) for c in clientes]

    async def count_clientes(self, filters: schemas.ClienteFilter) -> int:
        return await self.uow.clientes.count_search(
            nombre=filters.nombre,
            nit=filters.nit,
            id_tipo_cliente=filters.id_tipo_cliente,
            is_active=filters.is_active,
            fecha_desde=filters.fecha_desde,
            fecha_hasta=filters.fecha_hasta,
        )

    async def get_cliente(self, id_cliente: int) -> schemas.ClienteOut:
        cliente = await self.uow.clientes.get_by_id(id_cliente)
        if not cliente:
            raise HTTPException(status_code=404, detail="Cliente no encontrado")
        return self._to_out(cliente)

    async def get_by_nit(self, nit: str) -> schemas.ClienteOut:
        cliente = await self.uow.clientes.get_by_nit(nit)
        if not cliente:
            raise HTTPException(status_code=404, detail=f"Cliente con NIT '{nit}' no encontrado")
        return self._to_out(cliente)

    async def create_cliente(
        self, data: schemas.ClienteCreate, actor_id: int = 0
    ) -> schemas.ClienteOut:
        await self._validate_tipo_cliente(data.id_tipo_cliente)
        await self._validate_municipio(data.id_municipio)

        nit_clean = data.nit.strip().upper()
        if nit_clean != "CF":
            existing = await self.uow.clientes.get_by_nit(nit_clean)
            if existing and existing.estado == "Activo":
                raise HTTPException(status_code=400, detail="Ya existe un cliente con ese NIT")

        cliente = Cliente(
            id_tipo_cliente=data.id_tipo_cliente,
            id_municipio=data.id_municipio,
            nit=nit_clean,
            nombre_comercial=data.nombre_comercial.strip(),
            direccion=data.direccion.strip() if data.direccion else None,
            estado="Activo",
        )
        try:
            created = await self.uow.clientes.create(cliente)
            if data.contactos:
                for ct in data.contactos:
                    contacto = ClienteContacto(
                        id_cliente=created.id_cliente,
                        telefono=ct.telefono.strip() if ct.telefono else None,
                        correo=ct.correo.strip() if ct.correo else None,
                    )
                    await self.uow.clientes.add_contacto(contacto)
            await self.uow.commit()
        except Exception:
            await self.uow.rollback()
            raise

        fresh = await self.uow.clientes.get_by_id(created.id_cliente)
        audit_actor = actor_id or 1
        await log_action(
            self.uow.db,
            id_usuario=audit_actor,
            tabla="cliente",
            accion="CREAR",
            descripcion=f"Creó cliente '{fresh.nombre_comercial}' (NIT {fresh.nit}, ID {fresh.id_cliente}).",
        )
        await self.uow.commit()
        return self._to_out(fresh)

    async def update_cliente(
        self, id_cliente: int, data: schemas.ClienteUpdate, actor_id: int = 0
    ) -> schemas.ClienteOut:
        cliente = await self.uow.clientes.get_by_id(id_cliente)
        if not cliente:
            raise HTTPException(status_code=404, detail="Cliente no encontrado")

        if data.id_tipo_cliente is not None:
            await self._validate_tipo_cliente(data.id_tipo_cliente)
            cliente.id_tipo_cliente = data.id_tipo_cliente

        if data.id_municipio is not None:
            await self._validate_municipio(data.id_municipio)
            cliente.id_municipio = data.id_municipio

        if data.nit is not None:
            nit_clean = data.nit.strip().upper()
            if cliente.nit.strip().upper() == "CF" and nit_clean != "CF":
                raise HTTPException(
                    status_code=400,
                    detail="No se puede modificar el NIT del cliente Consumidor Final (CF)",
                )
            if cliente.nit.strip().upper() != "CF" and nit_clean == "CF":
                raise HTTPException(
                    status_code=400,
                    detail="No se puede asignar el NIT CF a otro cliente",
                )
            if nit_clean != "CF":
                duplicate = await self.uow.clientes.get_by_nit(nit_clean)
                if duplicate and duplicate.id_cliente != id_cliente and duplicate.estado == "Activo":
                    raise HTTPException(status_code=400, detail="Ya existe un cliente con ese NIT")
            cliente.nit = nit_clean

        if data.nombre_comercial is not None:
            cliente.nombre_comercial = data.nombre_comercial.strip()
        if data.direccion is not None:
            cliente.direccion = data.direccion.strip() if data.direccion else None
        if data.estado is not None:
            # Regla de negocio: cliente CF (id=1) no puede ponerse inactivo
            if cliente.nit.strip().upper() == "CF" and data.estado == "Inactivo":
                raise HTTPException(
                    status_code=400, detail="No se puede desactivar el cliente Consumidor Final (CF)"
                )
            cliente.estado = data.estado

        try:
            if data.contactos is not None:
                await self.uow.clientes.delete_contactos(id_cliente)
                for ct in data.contactos:
                    contacto = ClienteContacto(
                        id_cliente=id_cliente,
                        telefono=ct.telefono.strip() if ct.telefono else None,
                        correo=ct.correo.strip() if ct.correo else None,
                    )
                    await self.uow.clientes.add_contacto(contacto)

            await self.uow.clientes.update(cliente)
            await self.uow.commit()
        except Exception:
            await self.uow.rollback()
            raise

        fresh = await self.uow.clientes.get_by_id(id_cliente)
        audit_actor = actor_id or 1
        await log_action(
            self.uow.db,
            id_usuario=audit_actor,
            tabla="cliente",
            accion="EDITAR",
            descripcion=f"Editó cliente '{fresh.nombre_comercial}' (NIT {fresh.nit}, ID {fresh.id_cliente}).",
        )
        await self.uow.commit()
        return self._to_out(fresh)

    async def deactivate_cliente(self, id_cliente: int, actor_id: int = 0) -> schemas.ClienteOut:
        """Eliminación lógica (soft delete). Protege el registro de Consumidor Final."""
        cliente = await self.uow.clientes.get_by_id(id_cliente)
        if not cliente:
            raise HTTPException(status_code=404, detail="Cliente no encontrado")

        if cliente.nit.strip().upper() == "CF":
            raise HTTPException(
                status_code=400, detail="No se puede desactivar el cliente Consumidor Final (CF)"
            )

        cliente.estado = "Inactivo"
        try:
            await self.uow.clientes.update(cliente)
            await self.uow.commit()
        except Exception:
            await self.uow.rollback()
            raise

        fresh = await self.uow.clientes.get_by_id(id_cliente)
        audit_actor = actor_id or 1
        await log_action(
            self.uow.db,
            id_usuario=audit_actor,
            tabla="cliente",
            accion="DESACTIVAR",
            descripcion=f"Desactivó cliente '{fresh.nombre_comercial}' (NIT {fresh.nit}, ID {fresh.id_cliente}).",
        )
        await self.uow.commit()
        return self._to_out(fresh)

    # -----------------------------------------------------------------------
    # Exportación Excel / CSV
    # -----------------------------------------------------------------------

    _REPORT_HEADERS = ["ID", "NIT", "Nombre Comercial", "Tipo Cliente", "Municipio", "Dirección", "Estado"]

    def _cliente_row(self, c: schemas.ClienteOut) -> list:
        tipo = c.tipo_cliente.nombre if c.tipo_cliente else ""
        muni = c.municipio_nombre or ""
        return [c.id_cliente, c.nit, c.nombre_comercial, tipo, muni, c.direccion or "", c.estado]

    async def export_clientes(self, filters: schemas.ClienteFilter, fmt: str) -> StreamingResponse:
        clientes = await self.uow.clientes.search(
            nombre=filters.nombre,
            nit=filters.nit,
            id_tipo_cliente=filters.id_tipo_cliente,
            is_active=filters.is_active,
            fecha_desde=filters.fecha_desde,
            fecha_hasta=filters.fecha_hasta,
            sort_by=filters.sort_by,
            sort_order=filters.sort_order,
            limit=10_000,
            offset=0,
        )
        rows = [self._cliente_row(self._to_out(c)) for c in clientes]
        fecha_hoy = datetime.now().strftime("%Y-%m-%d")

        if fmt == "csv":
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(self._REPORT_HEADERS)
            writer.writerows(rows)
            output.seek(0)
            content = output.getvalue().encode("utf-8-sig")
            return StreamingResponse(
                iter([content]),
                media_type="text/csv; charset=utf-8",
                headers={"Content-Disposition": f'attachment; filename="reporte_clientes_{fecha_hoy}.csv"'},
            )

        wb = Workbook()
        ws = wb.active
        ws.title = "Clientes"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(fill_type="solid", fgColor="0F172A")
        header_align = Alignment(horizontal="center", vertical="center")

        for col_idx, header in enumerate(self._REPORT_HEADERS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        for row in rows:
            ws.append(row)

        col_widths = [8, 15, 35, 20, 20, 35, 12]
        for col_idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return StreamingResponse(
            iter([output.read()]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="reporte_clientes_{fecha_hoy}.xlsx"'},
        )
