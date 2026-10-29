import io
import csv
import openpyxl
from fastapi import HTTPException
from datetime import date, datetime, timezone
from typing import List, Optional
from decimal import Decimal

from app.schemas.reportes_schema import (
    GraficaResponse,
    KPITotalesDTO,
    ReporteConsolidadoResponse,
    TotalesConsolidadoDTO,
    ReporteExportadoDTO,
    VentaReporteDTO,
    MermaReporteDTO,
    CompraReporteDTO,
    UsuarioReporteDTO,
    ClienteReporteDTO,
    RankingItemDTO,
    RankingResponse,
)
from app.core.permissions import check_user_permission

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


class ReportesService:
    def __init__(self, uow):
        self.uow = uow

    async def _check_permission(self, user_id: int, permiso: str = "REPORTE_VER"):
        user = await self.uow.users.get_by_id(user_id)
        if not user:
            raise HTTPException(status_code=401, detail="Usuario no encontrado")

        allowed = await check_user_permission(self.uow.db, user, permiso)
        if not allowed:
            raise HTTPException(
                status_code=403,
                detail=f"Acceso denegado. Se requiere el permiso {permiso}.",
            )
        return user

    async def _check_admin(self, user_id: int):
        return await self._check_permission(user_id, "REPORTE_VER")

    # ---------------------------------------------------------------------------
    # Reporte Consolidado Operativo (Existente)
    # ---------------------------------------------------------------------------

    async def obtener_consolidado(self, fecha_inicio: date, fecha_fin: date, user_id: int) -> ReporteConsolidadoResponse:
        await self._check_permission(user_id)

        ventas_db = await self.uow.reportes.get_ventas_by_date(fecha_inicio, fecha_fin)
        mermas_db = await self.uow.reportes.get_mermas_by_date(fecha_inicio, fecha_fin)
        compras_db = await self.uow.reportes.get_compras_by_date(fecha_inicio, fecha_fin)

        ventas: List[VentaReporteDTO] = []
        total_ventas = Decimal("0.00")
        for v in ventas_db:
            vendedor = "Desconocido"
            if getattr(v, "usuario", None):
                vendedor = v.usuario.nombre_completo or v.usuario.username
            elif getattr(v, "id_usuario", None):
                vendedor = f"Usuario #{v.id_usuario}"

            # Construcción de resumen legible de artículos
            detalles_list = getattr(v, "detalles", []) or []
            item_strs = []
            for d in detalles_list:
                cant = getattr(d, "cantidad", 1)
                p_nom = "Artículo"
                if getattr(d, "variante", None) and getattr(d.variante, "producto", None):
                    p_nom = d.variante.producto.nombre
                    if getattr(d.variante, "color", None) or getattr(d.variante, "talla", None):
                        extra = " / ".join([x for x in (d.variante.color, d.variante.talla) if x])
                        p_nom = f"{p_nom} ({extra})"
                item_strs.append(f"{cant}x {p_nom}")
            articulos_resumen = ", ".join(item_strs) if item_strs else "Venta de mostrador"

            # Método de pago (Efectivo / Transferencia)
            pagos_list = getattr(v, "pagos", []) or []
            metodos = []
            for p in pagos_list:
                if getattr(p, "metodo_pago", None) and getattr(p.metodo_pago, "nombre", None):
                    metodos.append(p.metodo_pago.nombre)
            metodo_pago = ", ".join(set(metodos)) if metodos else "Efectivo"

            dt = v.fecha_venta
            if dt and dt.tzinfo is None:
                dt_aware = dt.replace(tzinfo=timezone.utc)
            else:
                dt_aware = dt

            subtotal = Decimal(str(v.subtotal or Decimal("0.00")))
            descuento = Decimal(str(v.descuento_total or Decimal("0.00")))
            total = Decimal(str(v.total_venta or Decimal("0.00")))
            total_ventas += total

            ventas.append(
                VentaReporteDTO(
                    id_venta=v.id_venta,
                    fecha=dt_aware,
                    fecha_venta=dt_aware,
                    estado=v.estado or "Completada",
                    vendedor=vendedor,
                    articulos=articulos_resumen,
                    metodo_pago=metodo_pago,
                    subtotal=subtotal,
                    descuento=descuento,
                    total=total,
                    total_venta=total,
                )
            )

        mermas: List[MermaReporteDTO] = []
        total_mermas = Decimal("0.00")
        for m in mermas_db:
            dt = m.fecha_merma
            if dt and dt.tzinfo is None:
                dt_aware = dt.replace(tzinfo=timezone.utc)
            else:
                dt_aware = dt

            costo_merma = Decimal("0.00")
            if getattr(m, "detalles", None):
                for d in m.detalles:
                    if getattr(d, "costo_perdida", None) is not None:
                        costo_merma += Decimal(str(d.costo_perdida))
            total_mermas += costo_merma

            mermas.append(
                MermaReporteDTO(
                    id_merma=m.id_merma,
                    fecha=dt_aware,
                    fecha_merma=dt_aware,
                    observaciones=m.observaciones,
                    total=costo_merma,
                )
            )

        compras: List[CompraReporteDTO] = []
        total_compras = Decimal("0.00")
        for c in compras_db:
            monto_compra = Decimal(str(c.total_compra or Decimal("0.00")))
            total_compras += monto_compra
            compras.append(
                CompraReporteDTO(
                    id_compra=c.id_compra,
                    fecha=c.fecha_ingreso,
                    fecha_ingreso=c.fecha_ingreso,
                    total=monto_compra,
                    total_compra=monto_compra,
                )
            )

        totales = TotalesConsolidadoDTO(
            ventas=total_ventas,
            mermas=total_mermas,
            compras=total_compras,
        )

        return ReporteConsolidadoResponse(
            totales=totales,
            ventas=ventas,
            mermas=mermas,
            compras=compras,
        )

    async def exportar_reporte(self, fecha_inicio: date, fecha_fin: date, formato: str, user_id: int) -> ReporteExportadoDTO:
        consolidado = await self.obtener_consolidado(fecha_inicio, fecha_fin, user_id)

        if formato.lower() == "excel":
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Reporte Consolidado"

            ws.append(["--- VENTAS ---"])
            ws.append(["ID", "Fecha Hora", "Total", "Estado"])
            for v in consolidado.ventas:
                dt = v.fecha_venta.replace(tzinfo=None) if v.fecha_venta.tzinfo else v.fecha_venta
                ws.append([v.id_venta, dt, float(v.total_venta), "Completada"])

            ws.append([])

            ws.append(["--- MERMAS ---"])
            ws.append(["ID", "Fecha Hora", "Observaciones"])
            for m in consolidado.mermas:
                dt = m.fecha_merma.replace(tzinfo=None) if m.fecha_merma.tzinfo else m.fecha_merma
                ws.append([m.id_merma, dt, m.observaciones or ""])

            ws.append([])

            ws.append(["--- COMPRAS ---"])
            ws.append(["ID", "Fecha Ingreso", "Total Compra"])
            for c in consolidado.compras:
                ws.append([c.id_compra, c.fecha_ingreso, float(c.total_compra)])

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            return ReporteExportadoDTO(
                filename=f"reporte_{fecha_inicio}_{fecha_fin}.xlsx",
                content=output.read(),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        elif formato.lower() == "csv":
            output = io.StringIO()
            writer = csv.writer(output)

            writer.writerow(["--- VENTAS ---"])
            writer.writerow(["ID", "Fecha Hora", "Total", "Estado"])
            for v in consolidado.ventas:
                writer.writerow([v.id_venta, v.fecha_venta, v.total_venta, "Completada"])

            writer.writerow([])
            writer.writerow(["--- MERMAS ---"])
            writer.writerow(["ID", "Fecha Hora", "Observaciones"])
            for m in consolidado.mermas:
                writer.writerow([m.id_merma, m.fecha_merma, m.observaciones or ""])

            writer.writerow([])
            writer.writerow(["--- COMPRAS ---"])
            writer.writerow(["ID", "Fecha Ingreso", "Total Compra"])
            for c in consolidado.compras:
                writer.writerow([c.id_compra, c.fecha_ingreso, c.total_compra])

            return ReporteExportadoDTO(
                filename=f"reporte_{fecha_inicio}_{fecha_fin}.csv",
                content=output.getvalue().encode("utf-8-sig"),
                media_type="text/csv",
            )
        elif formato.lower() == "pdf":
            buf = io.BytesIO()
            doc = SimpleDocTemplate(buf, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
            styles = getSampleStyleSheet()

            title_style = ParagraphStyle("CTitle", parent=styles["Heading1"], fontSize=15, leading=18, textColor=colors.HexColor("#0f172a"), alignment=1)
            sub_style = ParagraphStyle("CSub", parent=styles["Normal"], fontSize=9, leading=12, textColor=colors.HexColor("#64748b"), alignment=1, spaceAfter=8)
            cell_bold = ParagraphStyle("CBold", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.HexColor("#0f172a"))
            cell_text = ParagraphStyle("CText", parent=styles["Normal"], fontSize=7.5, leading=9, textColor=colors.HexColor("#334155"))

            elements = []
            elements.append(Paragraph("TIENDA EL REGALITO", title_style))
            elements.append(Paragraph("REPORTE OFICIAL - CONSOLIDADO OPERATIVO", sub_style))
            elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#0d7377"), spaceAfter=8))

            # Resumen KPI
            kpi_data = [
                [
                    Paragraph(f"<b>Desde:</b> {fecha_inicio}", cell_text),
                    Paragraph(f"<b>Hasta:</b> {fecha_fin}", cell_text),
                    Paragraph(f"<b>Total Ventas:</b> Q {float(consolidado.totales.ventas):,.2f}", cell_bold),
                    Paragraph(f"<b>Total Compras:</b> Q {float(consolidado.totales.compras):,.2f}", cell_text),
                ]
            ]
            t_kpi = Table(kpi_data, colWidths=[135, 135, 135, 135])
            t_kpi.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            elements.append(t_kpi)
            elements.append(Spacer(1, 10))

            # Ventas Recientes
            elements.append(Paragraph("<b>VENTAS REGISTRADAS EN EL PERÍODO</b>", cell_bold))
            elements.append(Spacer(1, 4))
            ventas_table = [
                [
                    Paragraph("<b>ID</b>", cell_bold),
                    Paragraph("<b>Fecha/Hora</b>", cell_bold),
                    Paragraph("<b>Artículos</b>", cell_bold),
                    Paragraph("<b>Vendedor</b>", cell_bold),
                    Paragraph("<b>Método Pago</b>", cell_bold),
                    Paragraph("<b>Total</b>", cell_bold),
                ]
            ]
            for v in consolidado.ventas:
                dt_str = v.fecha_venta.strftime("%d/%m/%Y %H:%M") if v.fecha_venta else "-"
                ventas_table.append([
                    Paragraph(str(v.id_venta), cell_text),
                    Paragraph(dt_str, cell_text),
                    Paragraph(v.articulos or "Venta general", cell_text),
                    Paragraph(v.vendedor or "-", cell_text),
                    Paragraph(v.metodo_pago or "Efectivo", cell_text),
                    Paragraph(f"Q {float(v.total_venta):,.2f}", cell_bold),
                ])
            t_v = Table(ventas_table, colWidths=[30, 80, 210, 80, 70, 70])
            t_v.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            elements.append(t_v)

            doc.build(elements)
            return ReporteExportadoDTO(
                filename=f"reporte_consolidado_{fecha_inicio}_{fecha_fin}.pdf",
                content=buf.getvalue(),
                media_type="application/pdf",
            )
        else:
            raise HTTPException(status_code=400, detail="Formato de exportación no soportado")

    # ---------------------------------------------------------------------------
    # Estadísticas del Dashboard de Reportes (KPI / Gráficas / Rankings)
    # ---------------------------------------------------------------------------

    async def obtener_kpi(self, fecha_inicio: date, fecha_fin: date, user_id: int) -> KPITotalesDTO:
        """Ensambla los totales agregados (SUM) del período seleccionado."""
        await self._check_permission(user_id)
        kpi = await self.uow.reportes.get_kpi_totales(fecha_inicio, fecha_fin)
        return KPITotalesDTO(**kpi)

    async def obtener_grafica(
        self,
        tipo: str,
        fecha_inicio: date,
        fecha_fin: date,
        user_id: int,
    ) -> GraficaResponse:
        """Devuelve los datos listos para Chart.js según la pestaña activa.

        - 'tendencia' -> barras: agrupadas por mes si el rango > 60 días (temporadas)
          o por día (DD/MM) si es menor.
        - 'categoria' -> barras horizontales: ingresos por categoría.
        - 'hora'      -> barras: nº de transacciones por franja horaria.
        """
        await self._check_permission(user_id)

        if tipo == "tendencia":
            puntos = await self.uow.reportes.get_tendencia_ventas(fecha_inicio, fecha_fin)
            dias_rango = (fecha_fin - fecha_inicio).days
            MESES_ES = [
                "", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
                "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"
            ]
            acum: dict = {}
            if dias_rango > 60:
                for dia, total in puntos:
                    clave = MESES_ES[dia.month] if fecha_inicio.year == fecha_fin.year else f"{MESES_ES[dia.month]} {dia.year}"
                    acum[clave] = acum.get(clave, Decimal("0.00")) + total
            else:
                for dia, total in puntos:
                    clave = dia.strftime("%d/%m")
                    acum[clave] = acum.get(clave, Decimal("0.00")) + total

            return GraficaResponse(
                tipo=tipo,
                etiquetas=list(acum.keys()),
                valores=list(acum.values()),
            )

        if tipo == "categoria":
            por_categoria = await self.uow.reportes.get_ventas_por_categoria(fecha_inicio, fecha_fin)
            return GraficaResponse(
                tipo=tipo,
                etiquetas=[nombre for nombre, _ in por_categoria],
                valores=[total for _, total in por_categoria],
            )

        if tipo == "hora":
            por_hora = await self.uow.reportes.get_ventas_por_hora(fecha_inicio, fecha_fin)
            return GraficaResponse(
                tipo=tipo,
                etiquetas=[f"{hora:02d}:00" for hora, _ in por_hora],
                valores=[Decimal(str(total)) for _, total in por_hora],
            )

        raise HTTPException(status_code=400, detail="Tipo de gráfica no válido")

    async def obtener_ranking(
        self,
        tipo_reporte: str,
        fecha_inicio: date,
        fecha_fin: date,
        limite: int,
        user_id: int,
    ) -> RankingResponse:
        """Widget de rankings (Top Productos / Clientes / Estancados).

        Consulta la base con ORDER BY total DESC LIMIT X y devuelve las filas
        listas para inyectar en el <ul> del frontend.
        """
        await self._check_permission(user_id)

        if tipo_reporte == "productos":
            filas = await self.uow.reportes.get_ranking_productos(fecha_inicio, fecha_fin, limite)
            items = [
                RankingItemDTO(nombre=nombre, valor=valor, cantidad=cantidad)
                for nombre, valor, cantidad in filas
            ]
        elif tipo_reporte == "clientes":
            filas = await self.uow.reportes.get_ranking_clientes(fecha_inicio, fecha_fin, limite)
            items = [
                RankingItemDTO(nombre=nombre, valor=valor, cantidad=cantidad)
                for nombre, valor, cantidad in filas
            ]
        elif tipo_reporte == "estancados":
            filas = await self.uow.reportes.get_ranking_estancados(fecha_inicio, fecha_fin, limite)
            items = [
                RankingItemDTO(nombre=nombre, valor=valor, cantidad=cantidad)
                for nombre, valor, cantidad in filas
            ]
        else:
            raise HTTPException(status_code=400, detail="Tipo de reporte no válido")

        return RankingResponse(tipo=tipo_reporte, items=items)

    # ---------------------------------------------------------------------------
    # Reporte de Usuarios
    # ---------------------------------------------------------------------------

    def _map_usuario(self, u, m: Optional[dict] = None) -> UsuarioReporteDTO:
        m = m or {}
        rol_nombre = u.rol.nombre if getattr(u, "rol", None) else "Sin rol"
        return UsuarioReporteDTO(
            id_usuario=u.id_usuario,
            nombre_completo=u.nombre_completo,
            username=u.username,
            id_rol=u.id_rol,
            rol_nombre=rol_nombre,
            estado=u.estado,
            fecha_creacion=u.fecha_creacion,
            ventas_count=m.get("ventas_count", 0),
            total_vendido=m.get("total_vendido", Decimal("0.00")),
            turnos_cerrados=m.get("turnos_cerrados", 0),
            diferencia_caja=m.get("diferencia_caja", Decimal("0.00")),
            ingresos_bodega=m.get("ingresos_bodega", 0),
        )

    async def obtener_reporte_usuarios(
        self,
        estado: Optional[str],
        fecha_desde: Optional[date],
        fecha_hasta: Optional[date],
        user_id: int,
    ) -> List[UsuarioReporteDTO]:
        await self._check_permission(user_id)
        usuarios = await self.uow.reportes.get_usuarios_reporte(
            estado=estado, fecha_desde=fecha_desde, fecha_hasta=fecha_hasta
        )
        metricas = await self.uow.reportes.get_usuarios_metricas_operativas(
            fecha_desde=fecha_desde, fecha_hasta=fecha_hasta
        )
        return [self._map_usuario(u, metricas.get(u.id_usuario)) for u in usuarios]

    async def obtener_ficha_usuario(self, id_usuario: int, user_id: int) -> UsuarioReporteDTO:
        await self._check_permission(user_id)
        u = await self.uow.reportes.get_usuario_by_id(id_usuario)
        if not u:
            raise HTTPException(status_code=404, detail="Usuario no encontrado")
        metricas = await self.uow.reportes.get_usuarios_metricas_operativas(None, None)
        return self._map_usuario(u, metricas.get(id_usuario))

    async def exportar_usuarios(
        self,
        estado: Optional[str],
        fecha_desde: Optional[date],
        fecha_hasta: Optional[date],
        formato: str,
        user_id: int,
    ) -> ReporteExportadoDTO:
        usuarios = await self.obtener_reporte_usuarios(estado, fecha_desde, fecha_hasta, user_id)
        fecha_str = datetime.now().strftime("%Y%m%d_%H%M")

        if formato.lower() == "excel":
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Rendimiento Usuarios"
            ws.append(["ID", "Usuario", "Nombre Completo", "Rol", "Ventas Realizadas", "Total Vendido (Q)", "Turnos Cerrados", "Diferencia Caja (Q)", "Ingresos Bodega", "Estado"])
            for u in usuarios:
                ws.append([
                    u.id_usuario,
                    u.username,
                    u.nombre_completo,
                    u.rol_nombre,
                    u.ventas_count,
                    float(u.total_vendido),
                    u.turnos_cerrados,
                    float(u.diferencia_caja),
                    u.ingresos_bodega,
                    u.estado,
                ])

            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            return ReporteExportadoDTO(
                filename=f"rendimiento_usuarios_{fecha_str}.xlsx",
                content=buf.read(),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        elif formato.lower() == "csv":
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(["ID", "Usuario", "Nombre Completo", "Rol", "Ventas Realizadas", "Total Vendido (Q)", "Turnos Cerrados", "Diferencia Caja (Q)", "Ingresos Bodega", "Estado"])
            for u in usuarios:
                writer.writerow([
                    u.id_usuario,
                    u.username,
                    u.nombre_completo,
                    u.rol_nombre,
                    u.ventas_count,
                    float(u.total_vendido),
                    u.turnos_cerrados,
                    float(u.diferencia_caja),
                    u.ingresos_bodega,
                    u.estado,
                ])

            return ReporteExportadoDTO(
                filename=f"rendimiento_usuarios_{fecha_str}.csv",
                content=output.getvalue().encode("utf-8-sig"),
                media_type="text/csv",
            )

        elif formato.lower() == "pdf":
            content = self._generar_pdf_usuarios_listado(usuarios, estado, fecha_desde, fecha_hasta)
            return ReporteExportadoDTO(
                filename=f"rendimiento_usuarios_{fecha_str}.pdf",
                content=content,
                media_type="application/pdf",
            )
        else:
            raise HTTPException(status_code=400, detail="Formato de exportación no soportado")

    async def exportar_ficha_usuario_pdf(self, id_usuario: int, user_id: int) -> ReporteExportadoDTO:
        usuario = await self.obtener_ficha_usuario(id_usuario, user_id)
        content = self._generar_pdf_ficha_usuario(usuario)
        return ReporteExportadoDTO(
            filename=f"ficha_usuario_{usuario.id_usuario}_{usuario.username}.pdf",
            content=content,
            media_type="application/pdf",
        )

    # ---------------------------------------------------------------------------
    # Reporte de Clientes
    # ---------------------------------------------------------------------------

    def _map_cliente(self, c) -> ClienteReporteDTO:
        tipo_nombre = c.tipo_cliente.nombre if getattr(c, "tipo_cliente", None) else "General"
        muni_nombre = c.municipio.nombre if getattr(c, "municipio", None) else None
        depto_nombre = None
        if getattr(c, "municipio", None) and getattr(c.municipio, "departamento", None):
            depto_nombre = c.municipio.departamento.nombre

        telefonos = [ct.telefono for ct in getattr(c, "contactos", []) if ct.telefono]
        correos = [ct.correo for ct in getattr(c, "contactos", []) if ct.correo]

        return ClienteReporteDTO(
            id_cliente=c.id_cliente,
            nit=c.nit,
            nombre_comercial=c.nombre_comercial,
            id_tipo_cliente=c.id_tipo_cliente,
            tipo_cliente_nombre=tipo_nombre,
            id_municipio=c.id_municipio,
            municipio_nombre=muni_nombre,
            departamento_nombre=depto_nombre,
            direccion=c.direccion,
            telefonos=telefonos,
            correos=correos,
            estado=c.estado,
            fecha_creacion=c.fecha_creacion,
        )

    async def obtener_reporte_clientes(
        self,
        estado: Optional[str],
        fecha_desde: Optional[date],
        fecha_hasta: Optional[date],
        user_id: int,
    ) -> List[ClienteReporteDTO]:
        await self._check_permission(user_id)
        clientes = await self.uow.reportes.get_clientes_reporte(
            estado=estado, fecha_desde=fecha_desde, fecha_hasta=fecha_hasta
        )
        return [self._map_cliente(c) for c in clientes]

    async def obtener_ficha_cliente(self, id_cliente: int, user_id: int) -> ClienteReporteDTO:
        await self._check_permission(user_id)
        c = await self.uow.reportes.get_cliente_by_id(id_cliente)
        if not c:
            raise HTTPException(status_code=404, detail="Cliente no encontrado")
        return self._map_cliente(c)

    async def exportar_clientes(
        self,
        estado: Optional[str],
        fecha_desde: Optional[date],
        fecha_hasta: Optional[date],
        formato: str,
        user_id: int,
    ) -> ReporteExportadoDTO:
        clientes = await self.obtener_reporte_clientes(estado, fecha_desde, fecha_hasta, user_id)
        fecha_str = datetime.now().strftime("%Y%m%d_%H%M")

        if formato.lower() == "excel":
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Reporte de Clientes"
            ws.append(["ID", "NIT", "Nombre Comercial", "Tipo", "Municipio", "Departamento", "Teléfonos", "Correos", "Estado", "Fecha Creación"])
            for c in clientes:
                fc = c.fecha_creacion.strftime("%Y-%m-%d %H:%M") if c.fecha_creacion else "-"
                ws.append([
                    c.id_cliente,
                    c.nit,
                    c.nombre_comercial,
                    c.tipo_cliente_nombre,
                    c.municipio_nombre or "-",
                    c.departamento_nombre or "-",
                    ", ".join(c.telefonos) if c.telefonos else "-",
                    ", ".join(c.correos) if c.correos else "-",
                    c.estado,
                    fc,
                ])

            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            return ReporteExportadoDTO(
                filename=f"reporte_clientes_{fecha_str}.xlsx",
                content=buf.read(),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        elif formato.lower() == "csv":
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(["ID", "NIT", "Nombre Comercial", "Tipo", "Municipio", "Departamento", "Telefonos", "Correos", "Estado", "Fecha Creacion"])
            for c in clientes:
                fc = c.fecha_creacion.strftime("%Y-%m-%d %H:%M") if c.fecha_creacion else "-"
                writer.writerow([
                    c.id_cliente,
                    c.nit,
                    c.nombre_comercial,
                    c.tipo_cliente_nombre,
                    c.municipio_nombre or "-",
                    c.departamento_nombre or "-",
                    "; ".join(c.telefonos) if c.telefonos else "-",
                    "; ".join(c.correos) if c.correos else "-",
                    c.estado,
                    fc,
                ])

            return ReporteExportadoDTO(
                filename=f"reporte_clientes_{fecha_str}.csv",
                content=output.getvalue().encode("utf-8-sig"),
                media_type="text/csv",
            )

        elif formato.lower() == "pdf":
            content = self._generar_pdf_clientes_listado(clientes, estado, fecha_desde, fecha_hasta)
            return ReporteExportadoDTO(
                filename=f"reporte_clientes_{fecha_str}.pdf",
                content=content,
                media_type="application/pdf",
            )
        else:
            raise HTTPException(status_code=400, detail="Formato de exportación no soportado")

    async def exportar_ficha_cliente_pdf(self, id_cliente: int, user_id: int) -> ReporteExportadoDTO:
        cliente = await self.obtener_ficha_cliente(id_cliente, user_id)
        content = self._generar_pdf_ficha_cliente(cliente)
        return ReporteExportadoDTO(
            filename=f"ficha_cliente_{cliente.id_cliente}_{cliente.nit}.pdf",
            content=content,
            media_type="application/pdf",
        )

    # ---------------------------------------------------------------------------
    # Generadores PDF con ReportLab
    # ---------------------------------------------------------------------------

    def _generar_pdf_usuarios_listado(
        self,
        usuarios: List[UsuarioReporteDTO],
        estado: Optional[str],
        fecha_desde: Optional[date],
        fecha_hasta: Optional[date],
    ) -> bytes:
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle("RTitle", parent=styles["Heading1"], fontSize=16, leading=20, textColor=colors.HexColor("#0f172a"), alignment=1)
        sub_style = ParagraphStyle("RSub", parent=styles["Normal"], fontSize=10, leading=13, textColor=colors.HexColor("#64748b"), alignment=1, spaceAfter=10)
        cell_bold = ParagraphStyle("CBold", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, leading=11, textColor=colors.HexColor("#0f172a"))
        cell_text = ParagraphStyle("CText", parent=styles["Normal"], fontSize=8, leading=10, textColor=colors.HexColor("#334155"))

        elements = []
        elements.append(Paragraph("TIENDA EL REGALITO", title_style))
        elements.append(Paragraph("REPORTE OPERATIVO DE RENDIMIENTO POR USUARIO / CAJERO", sub_style))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#0d7377"), spaceAfter=10))

        # Metadatos
        filtros_info = [
            [
                Paragraph("<b>Estado:</b> " + (estado or "Todos"), cell_text),
                Paragraph("<b>Fecha Desde:</b> " + (str(fecha_desde) if fecha_desde else "Inicio"), cell_text),
                Paragraph("<b>Fecha Hasta:</b> " + (str(fecha_hasta) if fecha_hasta else "Hoy"), cell_text),
                Paragraph("<b>Total Usuarios:</b> " + str(len(usuarios)), cell_text),
            ]
        ]
        t_meta = Table(filtros_info, colWidths=[135, 135, 135, 135])
        t_meta.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(t_meta)
        elements.append(Spacer(1, 10))

        # Tabla de rendimiento operativo
        table_data = [
            [
                Paragraph("<b>Usuario / Nombre</b>", cell_bold),
                Paragraph("<b>Rol</b>", cell_bold),
                Paragraph("<b>Ventas</b>", cell_bold),
                Paragraph("<b>Total (Q)</b>", cell_bold),
                Paragraph("<b>Turnos / Cuadre</b>", cell_bold),
                Paragraph("<b>Ingresos Bodega</b>", cell_bold),
            ]
        ]

        for u in usuarios:
            cuadre_str = f"{u.turnos_cerrados} turnos (Q {float(u.diferencia_caja):,.2f} dif)"
            table_data.append([
                Paragraph(f"<b>{u.nombre_completo}</b><br/>@{u.username}", cell_text),
                Paragraph(u.rol_nombre, cell_text),
                Paragraph(f"{u.ventas_count} ventas", cell_text),
                Paragraph(f"Q {float(u.total_vendido):,.2f}", cell_text),
                Paragraph(cuadre_str, cell_text),
                Paragraph(f"{u.ingresos_bodega} ingresos", cell_text),
            ])

        t_data = Table(table_data, colWidths=[140, 75, 75, 80, 100, 70])
        t_data.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elements.append(t_data)

        doc.build(elements)
        return buf.getvalue()

    def _generar_pdf_ficha_usuario(self, u: UsuarioReporteDTO) -> bytes:
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle("UTitle", parent=styles["Heading1"], fontSize=16, leading=20, textColor=colors.HexColor("#0f172a"), alignment=1)
        sub_style = ParagraphStyle("USub", parent=styles["Normal"], fontSize=11, leading=14, textColor=colors.HexColor("#64748b"), alignment=1, spaceAfter=14)
        label_style = ParagraphStyle("ULabel", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, textColor=colors.HexColor("#334155"))
        val_style = ParagraphStyle("UVal", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#0f172a"))

        elements = []
        elements.append(Paragraph("TIENDA EL REGALITO", title_style))
        elements.append(Paragraph(f"FICHA DE RENDIMIENTO OPERATIVO - {u.nombre_completo.upper()}", sub_style))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#0d7377"), spaceAfter=16))

        fc = u.fecha_creacion.strftime("%Y-%m-%d %H:%M") if u.fecha_creacion else "No registrada"
        data_card = [
            [Paragraph("ID Usuario:", label_style), Paragraph(str(u.id_usuario), val_style)],
            [Paragraph("Nombre Completo:", label_style), Paragraph(u.nombre_completo, val_style)],
            [Paragraph("Nombre de Usuario (Login):", label_style), Paragraph(f"@{u.username}", val_style)],
            [Paragraph("Rol Asignado:", label_style), Paragraph(u.rol_nombre, val_style)],
            [Paragraph("Estado Operativo:", label_style), Paragraph(u.estado, val_style)],
            [Paragraph("Ventas Realizadas:", label_style), Paragraph(f"{u.ventas_count} ventas registradas", val_style)],
            [Paragraph("Total Vendido:", label_style), Paragraph(f"Q {float(u.total_vendido):,.2f}", val_style)],
            [Paragraph("Turnos de Caja Cerrados:", label_style), Paragraph(f"{u.turnos_cerrados} turnos finalizados", val_style)],
            [Paragraph("Diferencia Acumulada de Caja:", label_style), Paragraph(f"Q {float(u.diferencia_caja):,.2f}", val_style)],
            [Paragraph("Ingresos de Mercadería (Bodega):", label_style), Paragraph(f"{u.ingresos_bodega} compras/ingresos procesados", val_style)],
            [Paragraph("Fecha de Registro:", label_style), Paragraph(fc, val_style)],
        ]

        t_card = Table(data_card, colWidths=[180, 350])
        t_card.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        elements.append(t_card)
        elements.append(Spacer(1, 30))

        # Firmas
        firma_data = [
            [
                Paragraph("___________________________________<br/><b>Firma del Usuario</b>", sub_style),
                Paragraph("___________________________________<br/><b>Firma Administrador(a)</b>", sub_style),
            ]
        ]
        t_firma = Table(firma_data, colWidths=[265, 265])
        t_firma.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER")]))
        elements.append(t_firma)

        doc.build(elements)
        return buf.getvalue()

    def _generar_pdf_clientes_listado(
        self,
        clientes: List[ClienteReporteDTO],
        estado: Optional[str],
        fecha_desde: Optional[date],
        fecha_hasta: Optional[date],
    ) -> bytes:
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle("CTitle", parent=styles["Heading1"], fontSize=16, leading=20, textColor=colors.HexColor("#0f172a"), alignment=1)
        sub_style = ParagraphStyle("CSub", parent=styles["Normal"], fontSize=10, leading=13, textColor=colors.HexColor("#64748b"), alignment=1, spaceAfter=10)
        cell_bold = ParagraphStyle("CBold", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.HexColor("#0f172a"))
        cell_text = ParagraphStyle("CText", parent=styles["Normal"], fontSize=8, leading=10, textColor=colors.HexColor("#334155"))

        elements = []
        elements.append(Paragraph("TIENDA EL REGALITO", title_style))
        elements.append(Paragraph("REPORTE OFICIAL DE CLIENTES", sub_style))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#2563eb"), spaceAfter=10))

        filtros_info = [
            [
                Paragraph("<b>Estado:</b> " + (estado or "Todos"), cell_text),
                Paragraph("<b>Fecha Desde:</b> " + (str(fecha_desde) if fecha_desde else "Inicio"), cell_text),
                Paragraph("<b>Fecha Hasta:</b> " + (str(fecha_hasta) if fecha_hasta else "Hoy"), cell_text),
                Paragraph("<b>Total Clientes:</b> " + str(len(clientes)), cell_text),
            ]
        ]
        t_meta = Table(filtros_info, colWidths=[135, 135, 135, 135])
        t_meta.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(t_meta)
        elements.append(Spacer(1, 10))

        table_data = [
            [
                Paragraph("<b>ID</b>", cell_bold),
                Paragraph("<b>NIT</b>", cell_bold),
                Paragraph("<b>Nombre Comercial / Razón Social</b>", cell_bold),
                Paragraph("<b>Tipo</b>", cell_bold),
                Paragraph("<b>Ubicación</b>", cell_bold),
                Paragraph("<b>Teléfono(s)</b>", cell_bold),
                Paragraph("<b>Estado</b>", cell_bold),
            ]
        ]

        for c in clientes:
            ubicacion = f"{c.municipio_nombre or ''}, {c.departamento_nombre or ''}".strip(", ") or "-"
            tels = ", ".join(c.telefonos) if c.telefonos else "-"
            table_data.append([
                Paragraph(str(c.id_cliente), cell_text),
                Paragraph(c.nit, cell_text),
                Paragraph(c.nombre_comercial, cell_text),
                Paragraph(c.tipo_cliente_nombre, cell_text),
                Paragraph(ubicacion, cell_text),
                Paragraph(tels, cell_text),
                Paragraph(c.estado, cell_text),
            ])

        t_data = Table(table_data, colWidths=[30, 65, 160, 75, 100, 60, 50])
        t_data.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(t_data)

        doc.build(elements)
        return buf.getvalue()

    def _generar_pdf_ficha_cliente(self, c: ClienteReporteDTO) -> bytes:
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle("FCTitle", parent=styles["Heading1"], fontSize=16, leading=20, textColor=colors.HexColor("#0f172a"), alignment=1)
        sub_style = ParagraphStyle("FCSub", parent=styles["Normal"], fontSize=11, leading=14, textColor=colors.HexColor("#64748b"), alignment=1, spaceAfter=14)
        label_style = ParagraphStyle("FCLabel", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, textColor=colors.HexColor("#334155"))
        val_style = ParagraphStyle("FCVal", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#0f172a"))

        elements = []
        elements.append(Paragraph("TIENDA EL REGALITO", title_style))
        elements.append(Paragraph(f"FICHA TÉCNICA DE CLIENTE #{c.id_cliente} - NIT: {c.nit}", sub_style))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#2563eb"), spaceAfter=16))

        fc = c.fecha_creacion.strftime("%Y-%m-%d %H:%M") if c.fecha_creacion else "No registrada"
        ubicacion = f"{c.municipio_nombre or ''}, {c.departamento_nombre or ''}".strip(", ") or "No asignada"
        tels = ", ".join(c.telefonos) if c.telefonos else "Ninguno"
        corrs = ", ".join(c.correos) if c.correos else "Ninguno"

        data_card = [
            [Paragraph("ID Cliente:", label_style), Paragraph(str(c.id_cliente), val_style)],
            [Paragraph("NIT:", label_style), Paragraph(c.nit, val_style)],
            [Paragraph("Nombre / Razón Social:", label_style), Paragraph(c.nombre_comercial, val_style)],
            [Paragraph("Tipo de Cliente:", label_style), Paragraph(c.tipo_cliente_nombre, val_style)],
            [Paragraph("Ubicación Geográfica:", label_style), Paragraph(ubicacion, val_style)],
            [Paragraph("Dirección Física:", label_style), Paragraph(c.direccion or "No registrada", val_style)],
            [Paragraph("Teléfono(s):", label_style), Paragraph(tels, val_style)],
            [Paragraph("Correo(s) Electrónico(s):", label_style), Paragraph(corrs, val_style)],
            [Paragraph("Estado:", label_style), Paragraph(c.estado, val_style)],
            [Paragraph("Fecha de Registro:", label_style), Paragraph(fc, val_style)],
        ]

        t_card = Table(data_card, colWidths=[180, 350])
        t_card.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        elements.append(t_card)

        doc.build(elements)
        return buf.getvalue()
