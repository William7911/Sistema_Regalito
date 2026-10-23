"""Servicio de catálogos generales (Fase 2 - Bloque 3).

Cubre el mantenimiento completo (CRUD + soft delete + reactivación + filtros +
ordenamiento + fichas + exportación Excel/CSV + auditoría + desactivación segura)
de: Motivos de merma, Métodos de pago, Proveedores, Departamentos geográficos y
Municipios (con dropdown dependiente departamento -> municipio).

Arquitectura DERCAS: controlador delgado -> este servicio (reglas de negocio) ->
repositorios (sin commit) -> el UnitOfWork (deps.py) hace commit/rollback.
"""
import csv
import io
from datetime import datetime
from typing import Optional, List
from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import (
    MotivoMerma, MetodoPago, Proveedor, DepartamentoGeografico, Municipio,
)
from app.schemas import catalogo_schema as schemas
from app.services.audit_service import log_action


class ConcreteCatalogoGeneralService:
    """Lógica de negocio para los catálogos del Bloque 3 de la Fase 2."""

    def __init__(self, uow):
        if isinstance(uow, AsyncSession):
            from app.api.deps import UnitOfWork
            self.uow = UnitOfWork(uow)
        else:
            self.uow = uow
        self.db = self.uow.db

    # -----------------------------------------------------------------------
    # Serializadores
    # -----------------------------------------------------------------------

    @staticmethod
    def _motivo_to_out(m: MotivoMerma) -> schemas.MotivoMermaOut:
        return schemas.MotivoMermaOut(id_motivo=m.id_motivo, descripcion=m.descripcion, estado=m.estado)

    @staticmethod
    def _metodo_to_out(m: MetodoPago) -> schemas.MetodoPagoOut:
        return schemas.MetodoPagoOut(id_metodo_pago=m.id_metodo_pago, nombre=m.nombre, estado=m.estado)

    @staticmethod
    def _proveedor_to_out(p: Proveedor) -> schemas.ProveedorOut:
        return schemas.ProveedorOut(
            id_proveedor=p.id_proveedor,
            nombre_contacto=p.nombre_contacto,
            distribuidora=p.distribuidora,
            telefono=p.telefono,
            correo=p.correo,
            cuenta_bancaria=p.cuenta_bancaria,
            direccion=p.direccion,
            estado=p.estado,
            fecha_creacion=p.fecha_creacion,
        )

    @staticmethod
    def _departamento_to_out(dep: DepartamentoGeografico) -> schemas.DepartamentoOut:
        munis = dep.__dict__.get("municipios")
        total = len(munis) if (munis and hasattr(munis, "__len__")) else 0
        return schemas.DepartamentoOut(
            id_departamento=dep.id_departamento,
            nombre=dep.nombre,
            estado=dep.estado,
            total_municipios=total,
        )

    @staticmethod
    def _municipio_to_out(m: Municipio) -> schemas.MunicipioOut:
        dep = m.__dict__.get("departamento")
        dep_nombre = dep.nombre if (dep and hasattr(dep, "nombre")) else None
        return schemas.MunicipioOut(
            id_municipio=m.id_municipio,
            id_departamento=m.id_departamento,
            nombre=m.nombre,
            estado=m.estado,
            departamento_nombre=dep_nombre,
        )

    # -----------------------------------------------------------------------
    # Helpers de exportación
    # -----------------------------------------------------------------------

    @staticmethod
    def _export(filename: str, headers: List[str], rows: List[list], fmt: str):
        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf)
            writer.writerow(headers)
            for r in rows:
                writer.writerow(r)
            buf.seek(0)
            return StreamingResponse(
                iter([buf.getvalue()]),
                media_type="text/csv",
                headers={"Content-Disposition": f'attachment; filename="{filename}.csv"'},
            )

        wb = Workbook()
        ws = wb.active
        ws.title = filename[:31]
        ws.append(headers)
        header_fill = PatternFill(start_color="0A5C5E", end_color="0A5C5E", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for r in rows:
            ws.append(r)
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'},
        )

    # =======================================================================
    # MOTIVOS DE MERMA
    # =======================================================================

    async def list_motivos_merma(self) -> List[schemas.MotivoMermaOut]:
        items = await self.uow.motivos_merma.list_all(only_active=True)
        return [self._motivo_to_out(m) for m in items]

    async def search_motivos_merma(self, filt: schemas.MotivoMermaFilter) -> List[schemas.MotivoMermaOut]:
        items = await self.uow.motivos_merma.search(
            search=filt.search,
            is_active=filt.is_active,
            sort_by=filt.sort_by or "descripcion",
            sort_order=filt.sort_order or "asc",
            limit=filt.limit,
            offset=filt.offset,
        )
        return [self._motivo_to_out(m) for m in items]

    async def count_motivos_merma(self, filt: schemas.MotivoMermaFilter) -> int:
        return await self.uow.motivos_merma.count_search(search=filt.search, is_active=filt.is_active)

    async def get_motivo_merma(self, id_motivo: int) -> schemas.MotivoMermaOut:
        m = await self.uow.motivos_merma.get_by_id(id_motivo)
        if not m:
            raise HTTPException(status_code=404, detail="Motivo de merma no encontrado")
        return self._motivo_to_out(m)

    async def create_motivo_merma(self, payload: schemas.MotivoMermaCreate, actor_id: int) -> schemas.MotivoMermaOut:
        descripcion = payload.descripcion.strip()
        existing = await self.uow.motivos_merma.get_by_descripcion(descripcion)
        if existing:
            raise HTTPException(status_code=400, detail="Ya existe un motivo de merma con esa descripción")
        obj = MotivoMerma(descripcion=descripcion, estado="Activo")
        created = await self.uow.motivos_merma.create(obj)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="CREATE", tabla="MotivoMerma",
            descripcion=f"Creó motivo de merma '{created.descripcion}' (ID: {created.id_motivo})",
        )
        await self.uow.commit()
        return self._motivo_to_out(created)

    async def update_motivo_merma(
        self, id_motivo: int, payload: schemas.MotivoMermaUpdate, actor_id: int
    ) -> schemas.MotivoMermaOut:
        m = await self.uow.motivos_merma.get_by_id(id_motivo)
        if not m:
            raise HTTPException(status_code=404, detail="Motivo de merma no encontrado")
        if payload.descripcion is not None:
            descripcion = payload.descripcion.strip()
            dup = await self.uow.motivos_merma.get_by_descripcion(descripcion)
            if dup and dup.id_motivo != id_motivo:
                raise HTTPException(status_code=400, detail="Ya existe un motivo de merma con esa descripción")
            m.descripcion = descripcion
        if payload.estado is not None:
            if payload.estado == "Inactivo":
                dependencias = await self.uow.motivos_merma.count_registros_asociados(id_motivo)
                if dependencias:
                    raise HTTPException(
                        status_code=400,
                        detail=f"No se puede desactivar el motivo de merma: hay {dependencias} registro(s) de merma asociado(s)",
                    )
            m.estado = payload.estado
        updated = await self.uow.motivos_merma.update(m)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="UPDATE", tabla="MotivoMerma",
            descripcion=f"Actualizó motivo de merma '{updated.descripcion}' (ID: {updated.id_motivo})",
        )
        await self.uow.commit()
        return self._motivo_to_out(updated)

    async def deactivate_motivo_merma(self, id_motivo: int, actor_id: int) -> schemas.MotivoMermaOut:
        m = await self.uow.motivos_merma.get_by_id(id_motivo)
        if not m:
            raise HTTPException(status_code=404, detail="Motivo de merma no encontrado")
        dependencias = await self.uow.motivos_merma.count_registros_asociados(id_motivo)
        if dependencias:
            raise HTTPException(
                status_code=400,
                detail=f"No se puede desactivar el motivo de merma: hay {dependencias} registro(s) de merma asociado(s)",
            )
        m.estado = "Inactivo"
        updated = await self.uow.motivos_merma.update(m)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="DEACTIVATE", tabla="MotivoMerma",
            descripcion=f"Desactivó motivo de merma '{updated.descripcion}' (ID: {updated.id_motivo})",
        )
        await self.uow.commit()
        return self._motivo_to_out(updated)

    async def activate_motivo_merma(self, id_motivo: int, actor_id: int) -> schemas.MotivoMermaOut:
        m = await self.uow.motivos_merma.get_by_id(id_motivo)
        if not m:
            raise HTTPException(status_code=404, detail="Motivo de merma no encontrado")
        m.estado = "Activo"
        updated = await self.uow.motivos_merma.update(m)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="ACTIVATE", tabla="MotivoMerma",
            descripcion=f"Reactivó motivo de merma '{updated.descripcion}' (ID: {updated.id_motivo})",
        )
        await self.uow.commit()
        return self._motivo_to_out(updated)

    async def export_motivos_merma(self, filt: schemas.MotivoMermaFilter, fmt: str):
        filt.limit = 5000
        filt.offset = 0
        items = await self.search_motivos_merma(filt)
        rows = [[m.id_motivo, m.descripcion, m.estado] for m in items]
        return self._export("motivos_merma", ["ID", "Descripción", "Estado"], rows, fmt)

    # =======================================================================
    # MÉTODOS DE PAGO
    # =======================================================================

    async def list_metodos_pago(self) -> List[schemas.MetodoPagoOut]:
        items = await self.uow.metodos_pago.list_all(only_active=True)
        return [self._metodo_to_out(m) for m in items]

    async def search_metodos_pago(self, filt: schemas.MetodoPagoFilter) -> List[schemas.MetodoPagoOut]:
        items = await self.uow.metodos_pago.search(
            search=filt.search,
            is_active=filt.is_active,
            sort_by=filt.sort_by or "nombre",
            sort_order=filt.sort_order or "asc",
            limit=filt.limit,
            offset=filt.offset,
        )
        return [self._metodo_to_out(m) for m in items]

    async def count_metodos_pago(self, filt: schemas.MetodoPagoFilter) -> int:
        return await self.uow.metodos_pago.count_search(search=filt.search, is_active=filt.is_active)

    async def get_metodo_pago(self, id_metodo: int) -> schemas.MetodoPagoOut:
        m = await self.uow.metodos_pago.get_by_id(id_metodo)
        if not m:
            raise HTTPException(status_code=404, detail="Método de pago no encontrado")
        return self._metodo_to_out(m)

    async def create_metodo_pago(self, payload: schemas.MetodoPagoCreate, actor_id: int) -> schemas.MetodoPagoOut:
        nombre = payload.nombre.strip()
        existing = await self.uow.metodos_pago.get_by_nombre(nombre)
        if existing:
            raise HTTPException(status_code=400, detail="Ya existe un método de pago con ese nombre")
        obj = MetodoPago(nombre=nombre, estado="Activo")
        created = await self.uow.metodos_pago.create(obj)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="CREATE", tabla="MetodoPago",
            descripcion=f"Creó método de pago '{created.nombre}' (ID: {created.id_metodo_pago})",
        )
        await self.uow.commit()
        return self._metodo_to_out(created)

    async def update_metodo_pago(
        self, id_metodo: int, payload: schemas.MetodoPagoUpdate, actor_id: int
    ) -> schemas.MetodoPagoOut:
        m = await self.uow.metodos_pago.get_by_id(id_metodo)
        if not m:
            raise HTTPException(status_code=404, detail="Método de pago no encontrado")
        if payload.nombre is not None:
            nombre = payload.nombre.strip()
            dup = await self.uow.metodos_pago.get_by_nombre(nombre)
            if dup and dup.id_metodo_pago != id_metodo:
                raise HTTPException(status_code=400, detail="Ya existe un método de pago con ese nombre")
            m.nombre = nombre
        if payload.estado is not None:
            if payload.estado == "Inactivo":
                dependencias = await self.uow.metodos_pago.count_pagos_ventas_completadas(id_metodo)
                if dependencias:
                    raise HTTPException(
                        status_code=400,
                        detail=f"No se puede desactivar el método de pago: hay {dependencias} pago(s) de ventas completadas asociado(s)",
                    )
            m.estado = payload.estado
        updated = await self.uow.metodos_pago.update(m)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="UPDATE", tabla="MetodoPago",
            descripcion=f"Actualizó método de pago '{updated.nombre}' (ID: {updated.id_metodo_pago})",
        )
        await self.uow.commit()
        return self._metodo_to_out(updated)

    async def deactivate_metodo_pago(self, id_metodo: int, actor_id: int) -> schemas.MetodoPagoOut:
        m = await self.uow.metodos_pago.get_by_id(id_metodo)
        if not m:
            raise HTTPException(status_code=404, detail="Método de pago no encontrado")
        dependencias = await self.uow.metodos_pago.count_pagos_ventas_completadas(id_metodo)
        if dependencias:
            raise HTTPException(
                status_code=400,
                detail=f"No se puede desactivar el método de pago: hay {dependencias} pago(s) de ventas completadas asociado(s)",
            )
        m.estado = "Inactivo"
        updated = await self.uow.metodos_pago.update(m)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="DEACTIVATE", tabla="MetodoPago",
            descripcion=f"Desactivó método de pago '{updated.nombre}' (ID: {updated.id_metodo_pago})",
        )
        await self.uow.commit()
        return self._metodo_to_out(updated)

    async def activate_metodo_pago(self, id_metodo: int, actor_id: int) -> schemas.MetodoPagoOut:
        m = await self.uow.metodos_pago.get_by_id(id_metodo)
        if not m:
            raise HTTPException(status_code=404, detail="Método de pago no encontrado")
        m.estado = "Activo"
        updated = await self.uow.metodos_pago.update(m)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="ACTIVATE", tabla="MetodoPago",
            descripcion=f"Reactivó método de pago '{updated.nombre}' (ID: {updated.id_metodo_pago})",
        )
        await self.uow.commit()
        return self._metodo_to_out(updated)

    async def export_metodos_pago(self, filt: schemas.MetodoPagoFilter, fmt: str):
        filt.limit = 5000
        filt.offset = 0
        items = await self.search_metodos_pago(filt)
        rows = [[m.id_metodo_pago, m.nombre, m.estado] for m in items]
        return self._export("metodos_pago", ["ID", "Nombre", "Estado"], rows, fmt)

    # =======================================================================
    # PROVEEDORES
    # =======================================================================

    async def list_proveedores(self) -> List[schemas.ProveedorOut]:
        items = await self.uow.proveedores.search(is_active=True, limit=1000)
        return [self._proveedor_to_out(p) for p in items]

    async def search_proveedores(self, filt: schemas.ProveedorFilter) -> List[schemas.ProveedorOut]:
        items = await self.uow.proveedores.search(
            search=filt.search,
            nombre=filt.nombre,
            distribuidora=filt.distribuidora,
            is_active=filt.is_active,
            fecha_desde=filt.fecha_desde,
            fecha_hasta=filt.fecha_hasta,
            sort_by=filt.sort_by or "distribuidora",
            sort_order=filt.sort_order or "asc",
            limit=filt.limit,
            offset=filt.offset,
        )
        return [self._proveedor_to_out(p) for p in items]

    async def count_proveedores(self, filt: schemas.ProveedorFilter) -> int:
        return await self.uow.proveedores.count_search(
            search=filt.search,
            nombre=filt.nombre,
            distribuidora=filt.distribuidora,
            is_active=filt.is_active,
            fecha_desde=filt.fecha_desde,
            fecha_hasta=filt.fecha_hasta,
        )

    async def get_proveedor(self, id_proveedor: int) -> schemas.ProveedorOut:
        p = await self.uow.proveedores.get_by_id(id_proveedor)
        if not p:
            raise HTTPException(status_code=404, detail="Proveedor no encontrado")
        return self._proveedor_to_out(p)

    async def create_proveedor(self, payload: schemas.ProveedorCreate, actor_id: int) -> schemas.ProveedorOut:
        distribuidora = payload.distribuidora.strip()
        existing = await self.uow.proveedores.get_by_distribuidora(distribuidora)
        if existing:
            raise HTTPException(status_code=400, detail="Ya existe un proveedor con esa distribuidora")
        obj = Proveedor(
            nombre_contacto=payload.nombre_contacto.strip(),
            distribuidora=distribuidora,
            telefono=payload.telefono.strip() if payload.telefono else None,
            correo=payload.correo.strip() if payload.correo else None,
            cuenta_bancaria=payload.cuenta_bancaria.strip() if payload.cuenta_bancaria else None,
            direccion=payload.direccion.strip() if payload.direccion else None,
            estado="Activo",
            fecha_creacion=datetime.utcnow(),
        )
        created = await self.uow.proveedores.create(obj)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="CREATE", tabla="Proveedor",
            descripcion=f"Creó proveedor '{created.distribuidora}' (ID: {created.id_proveedor})",
        )
        await self.uow.commit()
        return self._proveedor_to_out(created)

    async def update_proveedor(
        self, id_proveedor: int, payload: schemas.ProveedorUpdate, actor_id: int
    ) -> schemas.ProveedorOut:
        p = await self.uow.proveedores.get_by_id(id_proveedor)
        if not p:
            raise HTTPException(status_code=404, detail="Proveedor no encontrado")
        if payload.distribuidora is not None:
            distribuidora = payload.distribuidora.strip()
            if distribuidora.lower() != p.distribuidora.lower():
                dup = await self.uow.proveedores.get_by_distribuidora(distribuidora)
                if dup and dup.id_proveedor != id_proveedor:
                    raise HTTPException(status_code=400, detail="Ya existe un proveedor con esa distribuidora")
            p.distribuidora = distribuidora
        if payload.nombre_contacto is not None:
            p.nombre_contacto = payload.nombre_contacto.strip()
        if payload.telefono is not None:
            p.telefono = payload.telefono.strip() if payload.telefono else None
        if payload.correo is not None:
            p.correo = payload.correo.strip() if payload.correo else None
        if payload.cuenta_bancaria is not None:
            p.cuenta_bancaria = payload.cuenta_bancaria.strip() if payload.cuenta_bancaria else None
        if payload.direccion is not None:
            p.direccion = payload.direccion.strip() if payload.direccion else None
        if payload.estado is not None:
            if payload.estado == "Inactivo":
                dependencias = await self.uow.proveedores.count_ordenes_activas(id_proveedor)
                if dependencias:
                    raise HTTPException(
                        status_code=400,
                        detail=f"No se puede desactivar el proveedor: hay {dependencias} orden(es) de compra emitida(s) activa(s)",
                    )
            p.estado = payload.estado
        updated = await self.uow.proveedores.update(p)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="UPDATE", tabla="Proveedor",
            descripcion=f"Actualizó proveedor '{updated.distribuidora}' (ID: {updated.id_proveedor})",
        )
        await self.uow.commit()
        return self._proveedor_to_out(updated)

    async def deactivate_proveedor(self, id_proveedor: int, actor_id: int) -> schemas.ProveedorOut:
        p = await self.uow.proveedores.get_by_id(id_proveedor)
        if not p:
            raise HTTPException(status_code=404, detail="Proveedor no encontrado")
        dependencias = await self.uow.proveedores.count_ordenes_activas(id_proveedor)
        if dependencias:
            raise HTTPException(
                status_code=400,
                detail=f"No se puede desactivar el proveedor: hay {dependencias} orden(es) de compra emitida(s) activa(s)",
            )
        p.estado = "Inactivo"
        updated = await self.uow.proveedores.update(p)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="DEACTIVATE", tabla="Proveedor",
            descripcion=f"Desactivó proveedor '{updated.distribuidora}' (ID: {updated.id_proveedor})",
        )
        await self.uow.commit()
        return self._proveedor_to_out(updated)

    async def activate_proveedor(self, id_proveedor: int, actor_id: int) -> schemas.ProveedorOut:
        p = await self.uow.proveedores.get_by_id(id_proveedor)
        if not p:
            raise HTTPException(status_code=404, detail="Proveedor no encontrado")
        p.estado = "Activo"
        updated = await self.uow.proveedores.update(p)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="ACTIVATE", tabla="Proveedor",
            descripcion=f"Reactivó proveedor '{updated.distribuidora}' (ID: {updated.id_proveedor})",
        )
        await self.uow.commit()
        return self._proveedor_to_out(updated)

    async def export_proveedores(self, filt: schemas.ProveedorFilter, fmt: str):
        filt.limit = 5000
        filt.offset = 0
        items = await self.search_proveedores(filt)
        rows = [[
            p.id_proveedor,
            p.nombre_contacto,
            p.distribuidora,
            p.telefono or "",
            p.correo or "",
            p.cuenta_bancaria or "",
            p.direccion or "",
            p.estado,
            p.fecha_creacion.strftime("%Y-%m-%d %H:%M:%S") if p.fecha_creacion else "",
        ] for p in items]
        return self._export(
            "proveedores",
            ["ID", "Contacto", "Distribuidora", "Teléfono", "Correo", "Cuenta Bancaria", "Dirección", "Estado", "Fecha Creación"],
            rows,
            fmt,
        )

    # =======================================================================
    # DEPARTAMENTOS GEOGRÁFICOS
    # =======================================================================

    async def list_departamentos(self) -> List[schemas.DepartamentoOut]:
        items = await self.uow.departamentos_geo.list_all(only_active=True)
        return [self._departamento_to_out(d) for d in items]

    async def search_departamentos(self, filt: schemas.DepartamentoFilter) -> List[schemas.DepartamentoOut]:
        items = await self.uow.departamentos_geo.search(
            search=filt.search,
            is_active=filt.is_active,
            sort_by=filt.sort_by or "nombre",
            sort_order=filt.sort_order or "asc",
            limit=filt.limit,
            offset=filt.offset,
        )
        return [self._departamento_to_out(d) for d in items]

    async def count_departamentos(self, filt: schemas.DepartamentoFilter) -> int:
        return await self.uow.departamentos_geo.count_search(search=filt.search, is_active=filt.is_active)

    async def get_departamento(self, id_departamento: int) -> schemas.DepartamentoOut:
        dep = await self.uow.departamentos_geo.get_by_id(id_departamento)
        if not dep:
            raise HTTPException(status_code=404, detail="Departamento no encontrado")
        return self._departamento_to_out(dep)

    async def create_departamento(self, payload: schemas.DepartamentoCreate, actor_id: int) -> schemas.DepartamentoOut:
        nombre = payload.nombre.strip()
        existing = await self.uow.departamentos_geo.get_by_nombre(nombre)
        if existing:
            raise HTTPException(status_code=400, detail="Ya existe un departamento con ese nombre")
        obj = DepartamentoGeografico(nombre=nombre, estado="Activo")
        created = await self.uow.departamentos_geo.create(obj)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="CREATE", tabla="DepartamentoGeografico",
            descripcion=f"Creó departamento '{created.nombre}' (ID: {created.id_departamento})",
        )
        await self.uow.commit()
        return self._departamento_to_out(created)

    async def update_departamento(
        self, id_departamento: int, payload: schemas.DepartamentoUpdate, actor_id: int
    ) -> schemas.DepartamentoOut:
        dep = await self.uow.departamentos_geo.get_by_id(id_departamento)
        if not dep:
            raise HTTPException(status_code=404, detail="Departamento no encontrado")
        if payload.nombre is not None:
            nombre = payload.nombre.strip()
            if nombre.lower() != dep.nombre.lower():
                dup = await self.uow.departamentos_geo.get_by_nombre(nombre)
                if dup and dup.id_departamento != id_departamento:
                    raise HTTPException(status_code=400, detail="Ya existe un departamento con ese nombre")
            dep.nombre = nombre
        if payload.estado is not None:
            if payload.estado == "Inactivo":
                dependencias = await self.uow.departamentos_geo.count_municipios_activos(id_departamento)
                if dependencias:
                    raise HTTPException(
                        status_code=400,
                        detail=f"No se puede desactivar el departamento: hay {dependencias} municipio(s) activo(s) asociado(s)",
                    )
            dep.estado = payload.estado
        updated = await self.uow.departamentos_geo.update(dep)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="UPDATE", tabla="DepartamentoGeografico",
            descripcion=f"Actualizó departamento '{updated.nombre}' (ID: {updated.id_departamento})",
        )
        await self.uow.commit()
        return self._departamento_to_out(updated)

    async def deactivate_departamento(self, id_departamento: int, actor_id: int) -> schemas.DepartamentoOut:
        dep = await self.uow.departamentos_geo.get_by_id(id_departamento)
        if not dep:
            raise HTTPException(status_code=404, detail="Departamento no encontrado")
        dependencias = await self.uow.departamentos_geo.count_municipios_activos(id_departamento)
        if dependencias:
            raise HTTPException(
                status_code=400,
                detail=f"No se puede desactivar el departamento: hay {dependencias} municipio(s) activo(s) asociado(s)",
            )
        dep.estado = "Inactivo"
        updated = await self.uow.departamentos_geo.update(dep)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="DEACTIVATE", tabla="DepartamentoGeografico",
            descripcion=f"Desactivó departamento '{updated.nombre}' (ID: {updated.id_departamento})",
        )
        await self.uow.commit()
        return self._departamento_to_out(updated)

    async def activate_departamento(self, id_departamento: int, actor_id: int) -> schemas.DepartamentoOut:
        dep = await self.uow.departamentos_geo.get_by_id(id_departamento)
        if not dep:
            raise HTTPException(status_code=404, detail="Departamento no encontrado")
        dep.estado = "Activo"
        updated = await self.uow.departamentos_geo.update(dep)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="ACTIVATE", tabla="DepartamentoGeografico",
            descripcion=f"Reactivó departamento '{updated.nombre}' (ID: {updated.id_departamento})",
        )
        await self.uow.commit()
        return self._departamento_to_out(updated)

    async def export_departamentos(self, filt: schemas.DepartamentoFilter, fmt: str):
        filt.limit = 5000
        filt.offset = 0
        items = await self.search_departamentos(filt)
        rows = [[d.id_departamento, d.nombre, d.total_municipios or 0, d.estado] for d in items]
        return self._export("departamentos", ["ID", "Nombre", "N° Municipios", "Estado"], rows, fmt)

    # =======================================================================
    # MUNICIPIOS (dropdown dependiente departamento -> municipio)
    # =======================================================================

    async def list_municipios(self, id_departamento: Optional[int] = None) -> List[schemas.MunicipioOut]:
        items = await self.uow.municipios.list_by_departamento(id_departamento=id_departamento, only_active=True)
        return [self._municipio_to_out(m) for m in items]

    async def search_municipios(self, filt: schemas.MunicipioFilter) -> List[schemas.MunicipioOut]:
        items = await self.uow.municipios.search(
            search=filt.search,
            id_departamento=filt.id_departamento,
            is_active=filt.is_active,
            sort_by=filt.sort_by or "nombre",
            sort_order=filt.sort_order or "asc",
            limit=filt.limit,
            offset=filt.offset,
        )
        return [self._municipio_to_out(m) for m in items]

    async def count_municipios(self, filt: schemas.MunicipioFilter) -> int:
        return await self.uow.municipios.count_search(
            search=filt.search, id_departamento=filt.id_departamento, is_active=filt.is_active
        )

    async def get_municipio(self, id_municipio: int) -> schemas.MunicipioOut:
        m = await self.uow.municipios.get_by_id(id_municipio)
        if not m:
            raise HTTPException(status_code=404, detail="Municipio no encontrado")
        return self._municipio_to_out(m)

    async def create_municipio(self, payload: schemas.MunicipioCreate, actor_id: int) -> schemas.MunicipioOut:
        dep = await self.uow.departamentos_geo.get_by_id(payload.id_departamento)
        if not dep:
            raise HTTPException(status_code=400, detail="El departamento especificado no existe")
        nombre = payload.nombre.strip()
        dup = await self.uow.municipios.get_by_depto_nombre(payload.id_departamento, nombre)
        if dup:
            raise HTTPException(status_code=400, detail="Ya existe un municipio con ese nombre en este departamento")
        obj = Municipio(id_departamento=payload.id_departamento, nombre=nombre, estado="Activo")
        created = await self.uow.municipios.create(obj)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="CREATE", tabla="Municipio",
            descripcion=f"Creó municipio '{created.nombre}' en departamento {payload.id_departamento}",
        )
        await self.uow.commit()
        return self._municipio_to_out(created)

    async def update_municipio(
        self, id_municipio: int, payload: schemas.MunicipioUpdate, actor_id: int
    ) -> schemas.MunicipioOut:
        m = await self.uow.municipios.get_by_id(id_municipio)
        if not m:
            raise HTTPException(status_code=404, detail="Municipio no encontrado")
        if payload.nombre is not None:
            nombre = payload.nombre.strip()
            if nombre.lower() != m.nombre.lower():
                dup = await self.uow.municipios.get_by_depto_nombre(m.id_departamento, nombre)
                if dup and dup.id_municipio != id_municipio:
                    raise HTTPException(status_code=400, detail="Ya existe un municipio con ese nombre en este departamento")
            m.nombre = nombre
        if payload.estado is not None:
            if payload.estado == "Inactivo":
                dependencias = await self.uow.municipios.count_clientes_activos(id_municipio)
                if dependencias:
                    raise HTTPException(
                        status_code=400,
                        detail=f"No se puede desactivar el municipio: hay {dependencias} cliente(s) activo(s) asociado(s)",
                    )
            m.estado = payload.estado
        updated = await self.uow.municipios.update(m)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="UPDATE", tabla="Municipio",
            descripcion=f"Actualizó municipio '{updated.nombre}' (ID: {updated.id_municipio})",
        )
        await self.uow.commit()
        return self._municipio_to_out(updated)

    async def deactivate_municipio(self, id_municipio: int, actor_id: int) -> schemas.MunicipioOut:
        m = await self.uow.municipios.get_by_id(id_municipio)
        if not m:
            raise HTTPException(status_code=404, detail="Municipio no encontrado")
        dependencias = await self.uow.municipios.count_clientes_activos(id_municipio)
        if dependencias:
            raise HTTPException(
                status_code=400,
                detail=f"No se puede desactivar el municipio: hay {dependencias} cliente(s) activo(s) asociado(s)",
            )
        m.estado = "Inactivo"
        updated = await self.uow.municipios.update(m)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="DEACTIVATE", tabla="Municipio",
            descripcion=f"Desactivó municipio '{updated.nombre}' (ID: {updated.id_municipio})",
        )
        await self.uow.commit()
        return self._municipio_to_out(updated)

    async def activate_municipio(self, id_municipio: int, actor_id: int) -> schemas.MunicipioOut:
        m = await self.uow.municipios.get_by_id(id_municipio)
        if not m:
            raise HTTPException(status_code=404, detail="Municipio no encontrado")
        m.estado = "Activo"
        updated = await self.uow.municipios.update(m)
        await log_action(
            self.uow.db, id_usuario=actor_id, accion="ACTIVATE", tabla="Municipio",
            descripcion=f"Reactivó municipio '{updated.nombre}' (ID: {updated.id_municipio})",
        )
        await self.uow.commit()
        return self._municipio_to_out(updated)

    async def export_municipios(self, filt: schemas.MunicipioFilter, fmt: str):
        filt.limit = 5000
        filt.offset = 0
        items = await self.search_municipios(filt)
        rows = [[
            m.id_municipio,
            m.departamento_nombre or m.id_departamento,
            m.nombre,
            m.estado,
        ] for m in items]
        return self._export("municipios", ["ID", "Departamento", "Nombre", "Estado"], rows, fmt)