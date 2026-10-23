import csv
import io
from datetime import datetime
from typing import Optional, List
from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import Categoria, Subcategoria, UnidadMedida, Sucursal, AreaBodega
from app.interfaces.repositories import ProductRepository
from app.interfaces.services import CatalogoService
from app.schemas import catalogo_schema as schemas
from app.services.audit_service import log_action


class ConcreteCatalogoService(CatalogoService):
    """Lógica de negocio completa para los catálogos del Sistema Regalito (Fase 2 - Bloque 2):
    Sucursales, Áreas de bodega, Subcategorías y Unidades de medida.
    Soporta UnitOfWork, CRUD, soft delete, filtros, ordenamiento, exportación y compatibilidad con versiones previas."""

    def __init__(self, uow, product_repository: Optional[ProductRepository] = None):
        if isinstance(uow, AsyncSession):
            from app.api.deps import UnitOfWork
            self.uow = UnitOfWork(uow)
        else:
            self.uow = uow
        self.products = product_repository
        self.db = self.uow.db

    # -----------------------------------------------------------------------
    # Compatibilidad con métodos previos (dropdowns)
    # -----------------------------------------------------------------------

    async def list_subcategorias(self, include_inactive: bool = False) -> List[schemas.SubcategoriaOut]:
        if self.products:
            items = await self.products.list_subcategorias(include_inactive=include_inactive)
            return [self._subcategoria_to_out(s) for s in items]
        items = await self.uow.subcategorias.list_all(only_active=not include_inactive)
        return [self._subcategoria_to_out(s) for s in items]

    async def list_unidades(self) -> List[schemas.UnidadMedidaOut]:
        if self.products:
            items = await self.products.list_unidades()
            return [schemas.UnidadMedidaOut.model_validate(u) for u in items]
        items = await self.uow.unidades.list_all(only_active=True)
        return [schemas.UnidadMedidaOut.model_validate(u) for u in items]

    async def list_sucursales(self, include_inactive: bool = False) -> List[schemas.SucursalOut]:
        if self.products:
            items = await self.products.list_sucursales(include_inactive=include_inactive)
            return [self._sucursal_to_out(s) for s in items]
        items = await self.uow.sucursales.search(is_active=None if include_inactive else True, limit=500)
        return [self._sucursal_to_out(s) for s in items]

    async def list_areas(self, sucursal_id: Optional[int] = None) -> List[schemas.AreaBodegaOut]:
        if self.products:
            items = await self.products.list_areas(sucursal_id=sucursal_id)
            return [self._area_to_out(a) for a in items]
        items = await self.uow.areas.list_by_sucursal(id_sucursal=sucursal_id, only_active=True)
        return [self._area_to_out(a) for a in items]

    # -----------------------------------------------------------------------
    # Serializadores internos con relaciones
    # -----------------------------------------------------------------------

    def _sucursal_to_out(self, s: Sucursal) -> schemas.SucursalOut:
        areas = s.__dict__.get("areas")
        total_areas = len(areas) if (areas and hasattr(areas, "__len__")) else 0
        return schemas.SucursalOut(
            id_sucursal=s.id_sucursal,
            nombre=s.nombre,
            direccion=s.direccion,
            telefono=s.telefono,
            estado=s.estado,
            fecha_creacion=s.fecha_creacion,
            total_areas=total_areas,
        )

    def _area_to_out(self, a: AreaBodega) -> schemas.AreaBodegaOut:
        suc = a.__dict__.get("sucursal")
        suc_nombre = suc.nombre if (suc and hasattr(suc, "nombre")) else None
        return schemas.AreaBodegaOut(
            id_area=a.id_area,
            id_sucursal=a.id_sucursal,
            nombre_area=a.nombre_area,
            estado=a.estado,
            sucursal_nombre=suc_nombre,
        )

    def _subcategoria_to_out(self, s: Subcategoria) -> schemas.SubcategoriaOut:
        cat = s.__dict__.get("categoria")
        cat_nombre = cat.nombre if (cat and hasattr(cat, "nombre")) else None
        return schemas.SubcategoriaOut(
            id_subcategoria=s.id_subcategoria,
            id_categoria=s.id_categoria,
            nombre=s.nombre,
            estado=s.estado,
            categoria_nombre=cat_nombre,
        )

    # -----------------------------------------------------------------------
    # SUCURSAL CRUD
    # -----------------------------------------------------------------------

    async def create_sucursal(self, payload: schemas.SucursalCreate, actor_id: int) -> schemas.SucursalOut:
        nombre = payload.nombre.strip()
        existing = await self.uow.sucursales.get_by_nombre(nombre)
        if existing:
            raise HTTPException(status_code=400, detail="Ya existe una sucursal con ese nombre")

        suc = Sucursal(
            nombre=nombre,
            direccion=payload.direccion.strip(),
            telefono=payload.telefono.strip() if payload.telefono else None,
            estado="Activo",
            fecha_creacion=datetime.utcnow(),
        )
        created = await self.uow.sucursales.create(suc)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="CREAR",
            tabla="Sucursal",
            descripcion=f"Creó sucursal '{created.nombre}' (ID: {created.id_sucursal})",
        )
        await self.uow.commit()
        return self._sucursal_to_out(created)

    async def get_sucursal(self, id_sucursal: int) -> schemas.SucursalOut:
        suc = await self.uow.sucursales.get_by_id(id_sucursal)
        if not suc:
            raise HTTPException(status_code=404, detail="Sucursal no encontrada")
        return self._sucursal_to_out(suc)

    async def _validar_desactivacion_sucursal(self, id_sucursal: int) -> None:
        areas = await self.uow.sucursales.count_areas_activas(id_sucursal)
        if areas:
            raise HTTPException(
                status_code=400,
                detail=f"No se puede desactivar la sucursal: hay {areas} área(s) de bodega activa(s) asociada(s)",
            )
        inventario = await self.uow.sucursales.count_inventario_activo(id_sucursal)
        if inventario:
            raise HTTPException(
                status_code=400,
                detail=f"No se puede desactivar la sucursal: hay {inventario} registro(s) de inventario activo(s) asociado(s)",
            )

    async def update_sucursal(
        self, id_sucursal: int, payload: schemas.SucursalUpdate, actor_id: int
    ) -> schemas.SucursalOut:
        suc = await self.uow.sucursales.get_by_id(id_sucursal)
        if not suc:
            raise HTTPException(status_code=404, detail="Sucursal no encontrada")

        if payload.nombre is not None:
            nombre = payload.nombre.strip()
            if nombre.lower() != suc.nombre.lower():
                dup = await self.uow.sucursales.get_by_nombre(nombre)
                if dup and dup.id_sucursal != id_sucursal:
                    raise HTTPException(status_code=400, detail="Ya existe una sucursal con ese nombre")
            suc.nombre = nombre

        if payload.direccion is not None:
            suc.direccion = payload.direccion.strip()
        if payload.telefono is not None:
            suc.telefono = payload.telefono.strip() if payload.telefono else None
        if payload.estado is not None:
            if payload.estado == "Inactivo":
                await self._validar_desactivacion_sucursal(id_sucursal)
            suc.estado = payload.estado

        updated = await self.uow.sucursales.update(suc)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="EDITAR",
            tabla="Sucursal",
            descripcion=f"Actualizó sucursal '{updated.nombre}' (ID: {updated.id_sucursal})",
        )
        await self.uow.commit()
        return self._sucursal_to_out(updated)

    async def deactivate_sucursal(self, id_sucursal: int, actor_id: int) -> schemas.SucursalOut:
        suc = await self.uow.sucursales.get_by_id(id_sucursal)
        if not suc:
            raise HTTPException(status_code=404, detail="Sucursal no encontrada")
        await self._validar_desactivacion_sucursal(id_sucursal)
        suc.estado = "Inactivo"
        updated = await self.uow.sucursales.update(suc)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="DESACTIVAR",
            tabla="Sucursal",
            descripcion=f"Desactivó sucursal '{updated.nombre}' (ID: {updated.id_sucursal})",
        )
        await self.uow.commit()
        return self._sucursal_to_out(updated)

    async def activate_sucursal(self, id_sucursal: int, actor_id: int) -> schemas.SucursalOut:
        suc = await self.uow.sucursales.get_by_id(id_sucursal)
        if not suc:
            raise HTTPException(status_code=404, detail="Sucursal no encontrada")
        suc.estado = "Activo"
        updated = await self.uow.sucursales.update(suc)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="REACTIVAR",
            tabla="Sucursal",
            descripcion=f"Reactivó sucursal '{updated.nombre}' (ID: {updated.id_sucursal})",
        )
        await self.uow.commit()
        return self._sucursal_to_out(updated)

    async def search_sucursales(self, filt: schemas.SucursalFilter) -> List[schemas.SucursalOut]:
        items = await self.uow.sucursales.search(
            search=filt.search,
            is_active=filt.is_active,
            fecha_desde=filt.fecha_desde,
            fecha_hasta=filt.fecha_hasta,
            sort_by=filt.sort_by or "nombre",
            sort_order=filt.sort_order or "asc",
            limit=filt.limit,
            offset=filt.offset,
        )
        return [self._sucursal_to_out(s) for s in items]

    async def count_sucursales(self, filt: schemas.SucursalFilter) -> int:
        return await self.uow.sucursales.count(
            search=filt.search,
            is_active=filt.is_active,
            fecha_desde=filt.fecha_desde,
            fecha_hasta=filt.fecha_hasta,
        )

    async def export_sucursales(self, filt: schemas.SucursalFilter, fmt: str):
        filt.limit = 5000
        filt.offset = 0
        items = await self.search_sucursales(filt)

        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf)
            writer.writerow(["ID", "Nombre", "Dirección", "Teléfono", "Estado", "Fecha Creación"])
            for s in items:
                writer.writerow([
                    s.id_sucursal,
                    s.nombre,
                    s.direccion,
                    s.telefono or "",
                    s.estado,
                    s.fecha_creacion.strftime("%Y-%m-%d %H:%M:%S") if s.fecha_creacion else "",
                ])
            buf.seek(0)
            return StreamingResponse(
                iter([buf.getvalue()]),
                media_type="text/csv",
                headers={"Content-Disposition": 'attachment; filename="sucursales.csv"'},
            )

        wb = Workbook()
        ws = wb.active
        ws.title = "Sucursales"
        headers = ["ID", "Nombre", "Dirección", "Teléfono", "Estado", "Fecha Creación"]
        ws.append(headers)
        header_fill = PatternFill(start_color="0A5C5E", end_color="0A5C5E", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for s in items:
            ws.append([
                s.id_sucursal,
                s.nombre,
                s.direccion,
                s.telefono or "",
                s.estado,
                s.fecha_creacion.strftime("%Y-%m-%d %H:%M:%S") if s.fecha_creacion else "",
            ])

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="sucursales.xlsx"'},
        )

    # -----------------------------------------------------------------------
    # AREA BODEGA CRUD
    # -----------------------------------------------------------------------

    async def create_area_bodega(self, payload: schemas.AreaBodegaCreate, actor_id: int) -> schemas.AreaBodegaOut:
        suc = await self.uow.sucursales.get_by_id(payload.id_sucursal)
        if not suc:
            raise HTTPException(status_code=400, detail="La sucursal especificada no existe")

        nombre = payload.nombre_area.strip()
        dup = await self.uow.areas.get_by_nombre_sucursal(payload.id_sucursal, nombre)
        if dup:
            raise HTTPException(status_code=400, detail="Ya existe un área con ese nombre en esta sucursal")

        area = AreaBodega(
            id_sucursal=payload.id_sucursal,
            nombre_area=nombre,
            estado="Activo",
        )
        created = await self.uow.areas.create(area)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="CREAR",
            tabla="AreaBodega",
            descripcion=f"Creó área de bodega '{created.nombre_area}' en sucursal {payload.id_sucursal}",
        )
        await self.uow.commit()
        return self._area_to_out(created)

    async def get_area_bodega(self, id_area: int) -> schemas.AreaBodegaOut:
        area = await self.uow.areas.get_by_id(id_area)
        if not area:
            raise HTTPException(status_code=404, detail="Área de bodega no encontrada")
        return self._area_to_out(area)

    async def update_area_bodega(
        self, id_area: int, payload: schemas.AreaBodegaUpdate, actor_id: int
    ) -> schemas.AreaBodegaOut:
        area = await self.uow.areas.get_by_id(id_area)
        if not area:
            raise HTTPException(status_code=404, detail="Área de bodega no encontrada")

        if payload.nombre_area is not None:
            nombre = payload.nombre_area.strip()
            if nombre.lower() != area.nombre_area.lower():
                dup = await self.uow.areas.get_by_nombre_sucursal(area.id_sucursal, nombre)
                if dup and dup.id_area != id_area:
                    raise HTTPException(status_code=400, detail="Ya existe un área con ese nombre en esta sucursal")
            area.nombre_area = nombre

        if payload.estado is not None:
            if payload.estado == "Inactivo":
                dependencias = await self.uow.areas.count_inventario_activo(id_area)
                if dependencias:
                    raise HTTPException(
                        status_code=400,
                        detail=f"No se puede desactivar el área de bodega: hay {dependencias} registro(s) de inventario activo(s) asociado(s)",
                    )
            area.estado = payload.estado

        updated = await self.uow.areas.update(area)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="EDITAR",
            tabla="AreaBodega",
            descripcion=f"Actualizó área de bodega '{updated.nombre_area}' (ID: {updated.id_area})",
        )
        await self.uow.commit()
        return self._area_to_out(updated)

    async def deactivate_area_bodega(self, id_area: int, actor_id: int) -> schemas.AreaBodegaOut:
        area = await self.uow.areas.get_by_id(id_area)
        if not area:
            raise HTTPException(status_code=404, detail="Área de bodega no encontrada")
        dependencias = await self.uow.areas.count_inventario_activo(id_area)
        if dependencias:
            raise HTTPException(
                status_code=400,
                detail=f"No se puede desactivar el área de bodega: hay {dependencias} registro(s) de inventario activo(s) asociado(s)",
            )
        area.estado = "Inactivo"
        updated = await self.uow.areas.update(area)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="DESACTIVAR",
            tabla="AreaBodega",
            descripcion=f"Desactivó área de bodega '{updated.nombre_area}' (ID: {updated.id_area})",
        )
        await self.uow.commit()
        return self._area_to_out(updated)

    async def activate_area_bodega(self, id_area: int, actor_id: int) -> schemas.AreaBodegaOut:
        area = await self.uow.areas.get_by_id(id_area)
        if not area:
            raise HTTPException(status_code=404, detail="Área de bodega no encontrada")
        area.estado = "Activo"
        updated = await self.uow.areas.update(area)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="REACTIVAR",
            tabla="AreaBodega",
            descripcion=f"Reactivó área de bodega '{updated.nombre_area}' (ID: {updated.id_area})",
        )
        await self.uow.commit()
        return self._area_to_out(updated)

    async def search_areas(self, filt: schemas.AreaBodegaFilter) -> List[schemas.AreaBodegaOut]:
        items = await self.uow.areas.search(
            search=filt.search,
            id_sucursal=filt.id_sucursal,
            is_active=filt.is_active,
            sort_by=filt.sort_by or "nombre_area",
            sort_order=filt.sort_order or "asc",
            limit=filt.limit,
            offset=filt.offset,
        )
        return [self._area_to_out(a) for a in items]

    async def count_areas(self, filt: schemas.AreaBodegaFilter) -> int:
        return await self.uow.areas.count(
            search=filt.search,
            id_sucursal=filt.id_sucursal,
            is_active=filt.is_active,
        )

    async def export_areas(self, filt: schemas.AreaBodegaFilter, fmt: str):
        filt.limit = 5000
        filt.offset = 0
        items = await self.search_areas(filt)

        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf)
            writer.writerow(["ID", "Sucursal", "Nombre Área", "Estado"])
            for a in items:
                writer.writerow([a.id_area, a.sucursal_nombre or a.id_sucursal, a.nombre_area, a.estado])
            buf.seek(0)
            return StreamingResponse(
                iter([buf.getvalue()]),
                media_type="text/csv",
                headers={"Content-Disposition": 'attachment; filename="areas_bodega.csv"'},
            )

        wb = Workbook()
        ws = wb.active
        ws.title = "Areas"
        headers = ["ID", "Sucursal", "Nombre Área", "Estado"]
        ws.append(headers)
        header_fill = PatternFill(start_color="0A5C5E", end_color="0A5C5E", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for a in items:
            ws.append([a.id_area, a.sucursal_nombre or a.id_sucursal, a.nombre_area, a.estado])

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="areas_bodega.xlsx"'},
        )

    # -----------------------------------------------------------------------
    # SUBCATEGORIA CRUD
    # -----------------------------------------------------------------------

    async def create_subcategoria(self, payload: schemas.SubcategoriaCreate, actor_id: int) -> schemas.SubcategoriaOut:
        res = await self.uow.db.execute(select(Categoria).where(Categoria.id_categoria == payload.id_categoria))
        cat = res.scalar_one_or_none()
        if not cat:
            raise HTTPException(status_code=400, detail="La categoría especificada no existe")

        nombre = payload.nombre.strip()
        dup = await self.uow.subcategorias.get_by_cat_nombre(payload.id_categoria, nombre)
        if dup:
            raise HTTPException(status_code=400, detail="Ya existe una subcategoría con ese nombre en esta categoría")

        sub = Subcategoria(
            id_categoria=payload.id_categoria,
            nombre=nombre,
            estado="Activo",
        )
        created = await self.uow.subcategorias.create(sub)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="CREAR",
            tabla="Subcategoria",
            descripcion=f"Creó subcategoría '{created.nombre}' en categoría {payload.id_categoria}",
        )
        await self.uow.commit()
        return self._subcategoria_to_out(created)

    async def get_subcategoria(self, id_sub: int) -> schemas.SubcategoriaOut:
        sub = await self.uow.subcategorias.get_by_id(id_sub)
        if not sub:
            raise HTTPException(status_code=404, detail="Subcategoría no encontrada")
        return self._subcategoria_to_out(sub)

    async def update_subcategoria(
        self, id_sub: int, payload: schemas.SubcategoriaUpdate, actor_id: int
    ) -> schemas.SubcategoriaOut:
        sub = await self.uow.subcategorias.get_by_id(id_sub)
        if not sub:
            raise HTTPException(status_code=404, detail="Subcategoría no encontrada")

        if payload.nombre is not None:
            nombre = payload.nombre.strip()
            if nombre.lower() != sub.nombre.lower():
                dup = await self.uow.subcategorias.get_by_cat_nombre(sub.id_categoria, nombre)
                if dup and dup.id_subcategoria != id_sub:
                    raise HTTPException(status_code=400, detail="Ya existe una subcategoría con ese nombre en esta categoría")
            sub.nombre = nombre

        if payload.estado is not None:
            if payload.estado == "Inactivo":
                dependencias = await self.uow.subcategorias.count_productos_activos(id_sub)
                if dependencias:
                    raise HTTPException(
                        status_code=400,
                        detail=f"No se puede desactivar la subcategoría: hay {dependencias} producto(s) activo(s) asociado(s)",
                    )
            sub.estado = payload.estado

        updated = await self.uow.subcategorias.update(sub)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="EDITAR",
            tabla="Subcategoria",
            descripcion=f"Actualizó subcategoría '{updated.nombre}' (ID: {updated.id_subcategoria})",
        )
        await self.uow.commit()
        return self._subcategoria_to_out(updated)

    async def deactivate_subcategoria(self, id_sub: int, actor_id: int) -> schemas.SubcategoriaOut:
        sub = await self.uow.subcategorias.get_by_id(id_sub)
        if not sub:
            raise HTTPException(status_code=404, detail="Subcategoría no encontrada")
        dependencias = await self.uow.subcategorias.count_productos_activos(id_sub)
        if dependencias:
            raise HTTPException(
                status_code=400,
                detail=f"No se puede desactivar la subcategoría: hay {dependencias} producto(s) activo(s) asociado(s)",
            )
        sub.estado = "Inactivo"
        updated = await self.uow.subcategorias.update(sub)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="DESACTIVAR",
            tabla="Subcategoria",
            descripcion=f"Desactivó subcategoría '{updated.nombre}' (ID: {updated.id_subcategoria})",
        )
        await self.uow.commit()
        return self._subcategoria_to_out(updated)

    async def activate_subcategoria(self, id_sub: int, actor_id: int) -> schemas.SubcategoriaOut:
        sub = await self.uow.subcategorias.get_by_id(id_sub)
        if not sub:
            raise HTTPException(status_code=404, detail="Subcategoría no encontrada")
        sub.estado = "Activo"
        updated = await self.uow.subcategorias.update(sub)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="REACTIVAR",
            tabla="Subcategoria",
            descripcion=f"Reactivó subcategoría '{updated.nombre}' (ID: {updated.id_subcategoria})",
        )
        await self.uow.commit()
        return self._subcategoria_to_out(updated)

    async def search_subcategorias(self, filt: schemas.SubcategoriaFilter) -> List[schemas.SubcategoriaOut]:
        items = await self.uow.subcategorias.search(
            search=filt.search,
            id_categoria=filt.id_categoria,
            is_active=filt.is_active,
            sort_by=filt.sort_by or "nombre",
            sort_order=filt.sort_order or "asc",
            limit=filt.limit,
            offset=filt.offset,
        )
        return [self._subcategoria_to_out(s) for s in items]

    async def count_subcategorias(self, filt: schemas.SubcategoriaFilter) -> int:
        return await self.uow.subcategorias.count(
            search=filt.search,
            id_categoria=filt.id_categoria,
            is_active=filt.is_active,
        )

    async def export_subcategorias(self, filt: schemas.SubcategoriaFilter, fmt: str):
        filt.limit = 5000
        filt.offset = 0
        items = await self.search_subcategorias(filt)

        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf)
            writer.writerow(["ID", "Categoría", "Nombre Subcategoría", "Estado"])
            for s in items:
                writer.writerow([s.id_subcategoria, s.categoria_nombre or s.id_categoria, s.nombre, s.estado])
            buf.seek(0)
            return StreamingResponse(
                iter([buf.getvalue()]),
                media_type="text/csv",
                headers={"Content-Disposition": 'attachment; filename="subcategorias.csv"'},
            )

        wb = Workbook()
        ws = wb.active
        ws.title = "Subcategorias"
        headers = ["ID", "Categoría", "Nombre Subcategoría", "Estado"]
        ws.append(headers)
        header_fill = PatternFill(start_color="0A5C5E", end_color="0A5C5E", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for s in items:
            ws.append([s.id_subcategoria, s.categoria_nombre or s.id_categoria, s.nombre, s.estado])

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="subcategorias.xlsx"'},
        )

    # -----------------------------------------------------------------------
    # UNIDAD DE MEDIDA CRUD
    # -----------------------------------------------------------------------

    async def create_unidad_medida(self, payload: schemas.UnidadMedidaCreate, actor_id: int) -> schemas.UnidadMedidaOut:
        codigo = payload.codigo.strip().upper()
        existing = await self.uow.unidades.get_by_codigo(codigo)
        if existing:
            raise HTTPException(status_code=400, detail="Ya existe una unidad de medida con ese código")

        uni = UnidadMedida(
            codigo=codigo,
            descripcion=payload.descripcion.strip(),
            estado="Activo",
        )
        created = await self.uow.unidades.create(uni)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="CREAR",
            tabla="UnidadMedida",
            descripcion=f"Creó unidad de medida '{created.descripcion}' ({created.codigo})",
        )
        await self.uow.commit()
        return schemas.UnidadMedidaOut.model_validate(created)

    async def get_unidad_medida(self, id_unidad: int) -> schemas.UnidadMedidaOut:
        uni = await self.uow.unidades.get_by_id(id_unidad)
        if not uni:
            raise HTTPException(status_code=404, detail="Unidad de medida no encontrada")
        return schemas.UnidadMedidaOut.model_validate(uni)

    async def update_unidad_medida(
        self, id_unidad: int, payload: schemas.UnidadMedidaUpdate, actor_id: int
    ) -> schemas.UnidadMedidaOut:
        uni = await self.uow.unidades.get_by_id(id_unidad)
        if not uni:
            raise HTTPException(status_code=404, detail="Unidad de medida no encontrada")

        if payload.codigo is not None:
            codigo = payload.codigo.strip().upper()
            if codigo != uni.codigo:
                dup = await self.uow.unidades.get_by_codigo(codigo)
                if dup and dup.id_unidad != id_unidad:
                    raise HTTPException(status_code=400, detail="Ya existe una unidad de medida con ese código")
            uni.codigo = codigo

        if payload.descripcion is not None:
            uni.descripcion = payload.descripcion.strip()
        if payload.estado is not None:
            if payload.estado == "Inactivo":
                dependencias = await self.uow.unidades.count_productos_activos(id_unidad)
                if dependencias:
                    raise HTTPException(
                        status_code=400,
                        detail=f"No se puede desactivar la unidad de medida: hay {dependencias} producto(s) activo(s) asociado(s)",
                    )
            uni.estado = payload.estado

        updated = await self.uow.unidades.update(uni)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="EDITAR",
            tabla="UnidadMedida",
            descripcion=f"Actualizó unidad de medida '{updated.descripcion}' (ID: {updated.id_unidad})",
        )
        await self.uow.commit()
        return schemas.UnidadMedidaOut.model_validate(updated)

    async def deactivate_unidad_medida(self, id_unidad: int, actor_id: int) -> schemas.UnidadMedidaOut:
        uni = await self.uow.unidades.get_by_id(id_unidad)
        if not uni:
            raise HTTPException(status_code=404, detail="Unidad de medida no encontrada")
        dependencias = await self.uow.unidades.count_productos_activos(id_unidad)
        if dependencias:
            raise HTTPException(
                status_code=400,
                detail=f"No se puede desactivar la unidad de medida: hay {dependencias} producto(s) activo(s) asociado(s)",
            )
        uni.estado = "Inactivo"
        updated = await self.uow.unidades.update(uni)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="DESACTIVAR",
            tabla="UnidadMedida",
            descripcion=f"Desactivó unidad de medida '{updated.descripcion}' (ID: {updated.id_unidad})",
        )
        await self.uow.commit()
        return schemas.UnidadMedidaOut.model_validate(updated)

    async def activate_unidad_medida(self, id_unidad: int, actor_id: int) -> schemas.UnidadMedidaOut:
        uni = await self.uow.unidades.get_by_id(id_unidad)
        if not uni:
            raise HTTPException(status_code=404, detail="Unidad de medida no encontrada")
        uni.estado = "Activo"
        updated = await self.uow.unidades.update(uni)
        await log_action(
            self.uow.db,
            id_usuario=actor_id,
            accion="REACTIVAR",
            tabla="UnidadMedida",
            descripcion=f"Reactivó unidad de medida '{updated.descripcion}' (ID: {updated.id_unidad})",
        )
        await self.uow.commit()
        return schemas.UnidadMedidaOut.model_validate(updated)

    async def search_unidades(self, filt: schemas.UnidadMedidaFilter) -> List[schemas.UnidadMedidaOut]:
        items = await self.uow.unidades.search(
            search=filt.search,
            is_active=filt.is_active,
            sort_by=filt.sort_by or "descripcion",
            sort_order=filt.sort_order or "asc",
            limit=filt.limit,
            offset=filt.offset,
        )
        return [schemas.UnidadMedidaOut.model_validate(u) for u in items]

    async def count_unidades(self, filt: schemas.UnidadMedidaFilter) -> int:
        return await self.uow.unidades.count(
            search=filt.search,
            is_active=filt.is_active,
        )

    async def export_unidades(self, filt: schemas.UnidadMedidaFilter, fmt: str):
        filt.limit = 5000
        filt.offset = 0
        items = await self.search_unidades(filt)

        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf)
            writer.writerow(["ID", "Código", "Descripción", "Estado"])
            for u in items:
                writer.writerow([u.id_unidad, u.codigo, u.descripcion, u.estado])
            buf.seek(0)
            return StreamingResponse(
                iter([buf.getvalue()]),
                media_type="text/csv",
                headers={"Content-Disposition": 'attachment; filename="unidades_medida.csv"'},
            )

        wb = Workbook()
        ws = wb.active
        ws.title = "Unidades"
        headers = ["ID", "Código", "Descripción", "Estado"]
        ws.append(headers)
        header_fill = PatternFill(start_color="0A5C5E", end_color="0A5C5E", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for u in items:
            ws.append([u.id_unidad, u.codigo, u.descripcion, u.estado])

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="unidades_medida.xlsx"'},
        )
